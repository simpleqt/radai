"""
对话文件解析服务

三级解析策略：
  1. 纯文本 / Python 可直接读取的格式
  2. 富文档（PDF/DOCX/PPTX/XLSX）→ Python 库解析
  3. 解析失败 → 回退 MinerU OCR

图片类型（jpg/png/gif/webp）不做文本解析，直接标记为多模态。
"""

from __future__ import annotations

import asyncio
import csv
import re
import io
import json
import os
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional
from urllib.parse import urljoin

from backend.config.loggers import get_logger

logger = get_logger(__name__)

# ── 文件类型常量 ─────────────────────────────────────────

MULTIMODAL_EXTENSIONS: set[str] = {
    ".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp", ".svg",
}

PYTHON_PARSEABLE_EXTENSIONS: set[str] = {
    ".txt", ".md", ".csv", ".json", ".py", ".html", ".xml", ".log",
    ".yaml", ".yml", ".toml", ".ini", ".cfg", ".rst",
}

RICH_DOCUMENT_EXTENSIONS: set[str] = {
    ".pdf", ".docx", ".doc", ".pptx", ".ppt", ".xlsx", ".xls",
}

# MinerU 回退白名单（Python 解析失败时才会触发）
MINERU_FALLBACK_EXTENSIONS: set[str] = {".pdf", ".doc", ".ppt"}

# 短文档阈值（字符数）
SHORT_DOC_THRESHOLD = 5000

# 单文档文本截断上限
MAX_CONTENT_CHARS = 500_000


# ── 结果数据结构 ─────────────────────────────────────────

@dataclass
class ExtractedImage:
    """从文档中提取的图片"""
    data: bytes               # 图片二进制数据
    name: str                 # 图片文件名（如 image_0.png）
    caption: str = ""         # 图注 / 描述
    content_type: str = "image/png"  # MIME 类型


@dataclass
class ParseResult:
    """文件解析结果"""
    # multimodal / short_doc / long_doc / failed
    category: str
    # direct / python / mineru / none
    method: str
    content: Optional[str] = None
    content_length: int = 0
    error: Optional[str] = None
    # 给多模态文件使用：存储原始文件的 MIME type
    mime_type: Optional[str] = None
    # 从文档中提取的图片列表
    images: list[ExtractedImage] = field(default_factory=list)
    # 解析器已经拿到的远程图片 URL（例如 MinerU API 直接返回的图片）
    image_urls: list[dict] = field(default_factory=list)


# ── 主解析入口 ────────────────────────────────────────────

class ChatFileParser:
    """对话文件解析服务

    使用方式：
        parser = ChatFileParser()
        result = await parser.parse(file_bytes, filename, mime_type)
    """

    def __init__(self, mineru_api_url: Optional[str] = None):
        try:
            from backend.config.settings import settings
            self._mineru_api_url = mineru_api_url or getattr(settings, "mineru_api_url", None)
        except Exception:
            self._mineru_api_url = mineru_api_url

    async def parse(
        self,
        file_bytes: bytes,
        filename: str,
        mime_type: Optional[str] = None,
        parse_mode: str = "auto",
    ) -> ParseResult:
        """解析文件，返回类别和内容"""
        ext = Path(filename).suffix.lower()
        effective_parse_mode = str(parse_mode or "auto").strip().lower()

        # 图片 → 多模态
        if ext in MULTIMODAL_EXTENSIONS:
            return ParseResult(
                category="multimodal",
                method="none",
                content=None,
                mime_type=mime_type or _guess_mime(ext),
            )

        # 纯文本
        if ext in PYTHON_PARSEABLE_EXTENSIONS:
            content = await self._parse_text_file(file_bytes, ext)
            if content is not None:
                return self._route_by_length(content, method="direct")
            return ParseResult(
                category="failed",
                method="direct",
                error=f"文本文件读取失败: {filename}",
            )

        # 富文档
        if ext in RICH_DOCUMENT_EXTENSIONS:
            if effective_parse_mode == "mineru" and ext in MINERU_FALLBACK_EXTENSIONS | {".docx", ".pptx", ".xlsx"}:
                return await self._try_mineru(file_bytes, filename)
            return await self._parse_rich_document(file_bytes, filename, ext)

        # 未知格式：尝试 MinerU
        return await self._try_mineru(file_bytes, filename)

    # ── 纯文本解析 ────────────────────────────────────────

    async def _parse_text_file(
        self, data: bytes, ext: str
    ) -> Optional[str]:
        """读取纯文本内容（在线程池中执行以避免阻塞）"""
        return await asyncio.to_thread(self._parse_text_sync, data, ext)

    def _parse_text_sync(self, data: bytes, ext: str) -> Optional[str]:
        try:
            if ext == ".csv":
                text = data.decode("utf-8", errors="replace")
                reader = csv.reader(io.StringIO(text))
                lines = ["\t".join(row) for row in reader]
                return "\n".join(lines)
            if ext == ".json":
                text = data.decode("utf-8", errors="replace")
                obj = json.loads(text)
                return json.dumps(obj, ensure_ascii=False, indent=2)
            # 其他文本类型：直接解码
            return data.decode("utf-8", errors="replace")
        except Exception as e:
            logger.warning("[ChatFileParser] 文本解析失败 [%s]: %s", ext, e)
            return None

    # ── 富文档解析 ────────────────────────────────────────

    async def _parse_rich_document(
        self, data: bytes, filename: str, ext: str
    ) -> ParseResult:
        """富文档：先 Python 库解析，失败则 MinerU"""
        content = await asyncio.to_thread(self._python_parse_sync, data, ext)
        if content and _is_valid_content(content):
            logger.info(
                "[ChatFileParser] Python 解析成功: %s, 长度=%d", filename, len(content)
            )
            # 提取文档中的嵌入图片
            images = await asyncio.to_thread(self._extract_document_images, data, ext)
            if images:
                logger.info(
                    "[ChatFileParser] 提取到 %d 张图片: %s", len(images), filename
                )
            return self._route_by_length(content, method="python", images=images)

        # 回退 MinerU
        if ext in MINERU_FALLBACK_EXTENSIONS or ext in {".docx", ".pptx", ".xlsx"}:
            result = await self._try_mineru(data, filename)
            if result.category != "failed":
                return result

        return ParseResult(
            category="failed",
            method="python",
            error=f"无法解析文档内容，请确认文件格式正确: {filename}",
        )

    def _python_parse_sync(self, data: bytes, ext: str) -> Optional[str]:
        """Python 库解析（同步，在线程池中调用）"""
        try:
            if ext == ".pdf":
                return self._parse_pdf(data)
            elif ext == ".docx":
                return self._parse_docx(data)
            elif ext in (".xlsx", ".xls"):
                return self._parse_excel(data)
            elif ext in (".pptx", ".ppt"):
                return self._parse_pptx(data)
        except Exception as e:
            logger.warning("[ChatFileParser] Python 解析异常 [%s]: %s", ext, e)
        return None

    def _parse_pdf(self, data: bytes) -> Optional[str]:
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(data))
            parts = []
            for page in reader.pages:
                text = page.extract_text() or ""
                if text.strip():
                    parts.append(text.strip())
            return "\n\n".join(parts) if parts else None
        except ImportError:
            logger.debug("[ChatFileParser] pypdf 未安装，跳过 PDF Python 解析")
            return None

    def _parse_docx(self, data: bytes) -> Optional[str]:
        try:
            from docx import Document
            doc = Document(io.BytesIO(data))
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            # 提取表格
            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        paragraphs.append(" | ".join(cells))
            return "\n\n".join(paragraphs) if paragraphs else None
        except ImportError:
            logger.debug("[ChatFileParser] python-docx 未安装，跳过 DOCX 解析")
            return None

    def _parse_excel(self, data: bytes) -> Optional[str]:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            parts = []
            for sheet in wb.worksheets:
                parts.append(f"## Sheet: {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    cells = [str(c) if c is not None else "" for c in row]
                    if any(c.strip() for c in cells):
                        parts.append(" | ".join(cells))
            return "\n".join(parts) if parts else None
        except ImportError:
            logger.debug("[ChatFileParser] openpyxl 未安装，跳过 Excel 解析")
            return None

    def _parse_pptx(self, data: bytes) -> Optional[str]:
        try:
            from pptx import Presentation
            prs = Presentation(io.BytesIO(data))
            parts = []
            for i, slide in enumerate(prs.slides, 1):
                slide_texts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_texts.append(shape.text.strip())
                if slide_texts:
                    parts.append(f"## 幻灯片 {i}\n" + "\n".join(slide_texts))
            return "\n\n".join(parts) if parts else None
        except ImportError:
            logger.debug("[ChatFileParser] python-pptx 未安装，跳过 PPTX 解析")
            return None

    # ── 图片提取 ──────────────────────────────────────────

    def _extract_document_images(self, data: bytes, ext: str) -> list[ExtractedImage]:
        """从文档中提取嵌入图片（同步，在线程池中调用）"""
        try:
            if ext == ".pdf":
                return self._extract_pdf_images(data)
            elif ext == ".docx":
                return self._extract_docx_images(data)
            elif ext in (".pptx",):
                return self._extract_pptx_images(data)
        except Exception as e:
            logger.warning("[ChatFileParser] 图片提取异常 [%s]: %s", ext, e)
        return []

    def _extract_pdf_images(self, data: bytes) -> list[ExtractedImage]:
        """从 PDF 中提取嵌入图片"""
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(data))
            images: list[ExtractedImage] = []
            img_idx = 0
            for page_num, page in enumerate(reader.pages):
                try:
                    for image_key in page.images:
                        try:
                            img_data = image_key.data
                            if not img_data or len(img_data) < 100:
                                continue  # 跳过极小图片（可能是装饰元素）
                            img_name = f"image_{img_idx}.png"
                            # 检测 MIME 类型
                            ct = _detect_image_content_type(img_data)
                            caption = f"第{page_num + 1}页图片"
                            images.append(ExtractedImage(
                                data=img_data,
                                name=img_name,
                                caption=caption,
                                content_type=ct,
                            ))
                            img_idx += 1
                        except Exception as e:
                            logger.debug("[ChatFileParser] PDF 单张图片提取失败: %s", e)
                            continue
                except Exception:
                    continue
            return images
        except ImportError:
            logger.debug("[ChatFileParser] pypdf 未安装，跳过 PDF 图片提取")
            return []

    def _extract_docx_images(self, data: bytes) -> list[ExtractedImage]:
        """从 DOCX 中提取嵌入图片"""
        try:
            from docx import Document
            from docx.opc.constants import RELATIONSHIP_TYPE as RT
            doc = Document(io.BytesIO(data))
            images: list[ExtractedImage] = []
            img_idx = 0
            # 遍历文档中的所有关联部件（图片）
            for rel in doc.part.rels.values():
                if "image" in rel.reltype:
                    try:
                        img_blob = rel.target_part.blob
                        if not img_blob or len(img_blob) < 100:
                            continue
                        ct = rel.target_part.content_type or _detect_image_content_type(img_blob)
                        ext_map = {
                            "image/png": ".png",
                            "image/jpeg": ".jpg",
                            "image/gif": ".gif",
                            "image/bmp": ".bmp",
                        }
                        suffix = ext_map.get(ct, ".png")
                        img_name = f"image_{img_idx}{suffix}"
                        images.append(ExtractedImage(
                            data=img_blob,
                            name=img_name,
                            caption=f"文档图片 {img_idx + 1}",
                            content_type=ct,
                        ))
                        img_idx += 1
                    except Exception as e:
                        logger.debug("[ChatFileParser] DOCX 单张图片提取失败: %s", e)
                        continue
            return images
        except ImportError:
            logger.debug("[ChatFileParser] python-docx 未安装，跳过 DOCX 图片提取")
            return []

    def _extract_pptx_images(self, data: bytes) -> list[ExtractedImage]:
        """从 PPTX 中提取嵌入图片"""
        try:
            from pptx import Presentation
            from pptx.enum.shapes import MSO_SHAPE_TYPE
            prs = Presentation(io.BytesIO(data))
            images: list[ExtractedImage] = []
            img_idx = 0
            for slide_num, slide in enumerate(prs.slides, 1):
                for shape in slide.shapes:
                    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                        try:
                            image = shape.image
                            img_blob = image.blob
                            if not img_blob or len(img_blob) < 100:
                                continue
                            ct = image.content_type or "image/png"
                            ext_map = {
                                "image/png": ".png",
                                "image/jpeg": ".jpg",
                                "image/gif": ".gif",
                            }
                            suffix = ext_map.get(ct, ".png")
                            img_name = f"image_{img_idx}{suffix}"
                            images.append(ExtractedImage(
                                data=img_blob,
                                name=img_name,
                                caption=f"幻灯片{slide_num}图片",
                                content_type=ct,
                            ))
                            img_idx += 1
                        except Exception as e:
                            logger.debug("[ChatFileParser] PPTX 单张图片提取失败: %s", e)
                            continue
            return images
        except ImportError:
            logger.debug("[ChatFileParser] python-pptx 未安装，跳过 PPTX 图片提取")
            return []

    # ── MinerU 回退 ───────────────────────────────────────

    async def _try_mineru(self, data: bytes, filename: str) -> ParseResult:
        """调用 MinerU OCR 解析，并尝试提取图片"""
        if not self._mineru_api_url:
            logger.warning("[ChatFileParser] MinerU API URL 未配置，跳过 OCR 解析")
            return ParseResult(
                category="failed",
                method="mineru",
                error="MinerU 服务未配置",
            )

        tmp_path = None
        try:
            from backend.core.tools.ocr.mineru_ocr import MinerUOCR

            with tempfile.NamedTemporaryFile(
                suffix=Path(filename).suffix, delete=False
            ) as tmp:
                tmp.write(data)
                tmp_path = tmp.name

            ocr = MinerUOCR(api_url=self._mineru_api_url)
            try:
                # 提交并等待完成
                task_id = await ocr.submit_task_async(tmp_path)
                result = await ocr.wait_for_completion_async(task_id, timeout=300)
                if not result.is_success():
                    return ParseResult(
                        category="failed",
                        method="mineru",
                        error=f"MinerU 解析失败: {result.error_message or 'unknown'}",
                    )

                # 获取详细数据，包括图片
                detail = await ocr.get_task_data_async(
                    task_id,
                    include_fields=["md", "content_list", "images"],
                    upload_images=True,
                )

                content = detail.markdown or result.markdown
                if not content or not _is_valid_content(content):
                    return ParseResult(
                        category="failed",
                        method="mineru",
                        error=f"MinerU 解析未返回有效内容: {filename}",
                    )

                # 从 MinerU 结果中提取图片
                images, remote_image_urls = self._extract_mineru_images(
                    detail,
                    content,
                    markdown_url=getattr(detail, "md_url", None),
                )

                logger.info(
                    "[ChatFileParser] MinerU 解析成功: %s, 长度=%d, 图片=%d, 远程图片=%d",
                    filename, len(content), len(images), len(remote_image_urls),
                )
                result = self._route_by_length(content, method="mineru", images=images)
                result.image_urls = remote_image_urls
                return result
            finally:
                await ocr.close()

        except Exception as e:
            logger.error("[ChatFileParser] MinerU 解析异常: %s, error=%s", filename, e)
            return ParseResult(
                category="failed",
                method="mineru",
                error=str(e),
            )
        finally:
            if tmp_path and os.path.exists(tmp_path):
                try:
                    os.unlink(tmp_path)
                except Exception:
                    pass

    def _extract_mineru_images(
        self,
        ocr_result,
        markdown_content: str,
        markdown_url: str | None = None,
    ) -> tuple[list[ExtractedImage], list[dict]]:
        """从 MinerU OCR 结果中提取图片数据与远程图片 URL"""
        images: list[ExtractedImage] = []
        remote_image_urls: list[dict] = []
        seen_remote_keys: set[tuple[str, str]] = set()

        def _resolve_remote_path(raw_path: str) -> str:
            path = str(raw_path or "").strip()
            if not path:
                return ""
            if path.startswith(("http://", "https://")):
                return path
            if markdown_url and path:
                return urljoin(markdown_url, path)
            return ""

        def _normalize_caption(value, fallback: str) -> str:
            if isinstance(value, str):
                return value.strip() or fallback
            if isinstance(value, list):
                joined = " ".join(str(item).strip() for item in value if str(item).strip())
                return joined or fallback
            return fallback

        def _append_remote_image(*, raw_url: str, name: str, caption: str, content_type: str | None = None) -> None:
            resolved_url = _resolve_remote_path(raw_url)
            if not resolved_url:
                return
            safe_name = Path(str(name or "")).name or f"mineru_image_{len(remote_image_urls) + 1}.png"
            dedupe_key = (resolved_url, safe_name)
            if dedupe_key in seen_remote_keys:
                return
            seen_remote_keys.add(dedupe_key)
            payload = {
                "url": resolved_url,
                "name": safe_name,
                "caption": caption or f"文档图片 {len(remote_image_urls) + 1}",
            }
            if content_type:
                payload["content_type"] = content_type
            remote_image_urls.append(payload)

        try:
            # MinerU 返回的 images 字段是 dict: {filename: base64_or_url}
            mineru_images = getattr(ocr_result, 'images', None)
            if mineru_images and isinstance(mineru_images, dict):
                img_idx = 0
                for img_name, img_value in mineru_images.items():
                    try:
                        if isinstance(img_value, str):
                            if img_value.startswith(('http://', 'https://')):
                                _append_remote_image(
                                    raw_url=img_value,
                                    name=os.path.basename(img_name) or img_name,
                                    caption=f"文档图片 {img_idx + 1}",
                                )
                                img_idx += 1
                                continue
                            # base64 编码的图片
                            import base64
                            img_data = base64.b64decode(img_value)
                        elif isinstance(img_value, dict):
                            raw_url = (
                                img_value.get("url")
                                or img_value.get("img_url")
                                or img_value.get("image_url")
                                or img_value.get("img_path")
                            )
                            if raw_url:
                                _append_remote_image(
                                    raw_url=str(raw_url),
                                    name=img_value.get("name") or img_name,
                                    caption=_normalize_caption(
                                        img_value.get("caption") or img_value.get("image_caption"),
                                        f"文档图片 {img_idx + 1}",
                                    ),
                                    content_type=img_value.get("content_type"),
                                )
                                img_idx += 1
                                continue
                            raw_base64 = img_value.get("base64") or img_value.get("content")
                            if not raw_base64:
                                continue
                            import base64
                            img_data = base64.b64decode(raw_base64)
                        elif isinstance(img_value, bytes):
                            img_data = img_value
                        else:
                            continue

                        if len(img_data) < 100:
                            continue

                        ct = _detect_image_content_type(img_data)
                        images.append(ExtractedImage(
                            data=img_data,
                            name=img_name or f"mineru_image_{img_idx}.png",
                            caption=f"文档图片 {img_idx + 1}",
                            content_type=ct,
                        ))
                        img_idx += 1
                    except Exception as e:
                        logger.debug("[ChatFileParser] MinerU 单张图片提取失败: %s", e)
                        continue

            # 从 content_list 中提取（如果有结构化数据）
            content_list = getattr(ocr_result, 'content_list', None)
            if content_list and isinstance(content_list, list):
                img_idx = 0
                for item in content_list:
                    if isinstance(item, dict) and item.get('type') == 'image':
                        img_path = (
                            item.get('img_path')
                            or item.get('img_url')
                            or item.get('image_url')
                            or ""
                        )
                        caption = _normalize_caption(
                            item.get('image_caption') or item.get('img_caption') or item.get('caption'),
                            f"文档图片 {img_idx + 1}",
                        )
                        resolved_remote = _resolve_remote_path(str(img_path))
                        if resolved_remote:
                            _append_remote_image(
                                raw_url=resolved_remote,
                                name=os.path.basename(str(img_path)) or f"mineru_image_{img_idx + 1}.png",
                                caption=caption,
                            )
                            img_idx += 1
                            continue
                        if img_path and os.path.exists(img_path):
                            try:
                                with open(img_path, 'rb') as f:
                                    img_data = f.read()
                                if len(img_data) >= 100:
                                    ct = _detect_image_content_type(img_data)
                                    images.append(ExtractedImage(
                                        data=img_data,
                                        name=os.path.basename(img_path),
                                        caption=caption,
                                        content_type=ct,
                                    ))
                                    img_idx += 1
                            except Exception as e:
                                logger.debug("[ChatFileParser] MinerU content_list 图片读取失败: %s", e)

            # 最后从 markdown 中兜底提取远程图片引用，避免仅有 md 图片链接时漏掉
            for idx, match in enumerate(
                re.finditer(r"!\[([^\]]*)\]\(([^)]+)\)", markdown_content or ""),
                start=1,
            ):
                caption = (match.group(1) or "").strip() or f"文档图片 {idx}"
                raw_src = (match.group(2) or "").strip()
                _append_remote_image(
                    raw_url=raw_src,
                    name=os.path.basename(raw_src.split("?", 1)[0]) or f"markdown_image_{idx}.png",
                    caption=caption,
                )
        except Exception as e:
            logger.warning("[ChatFileParser] MinerU 图片提取整体异常: %s", e)
        return images, remote_image_urls

    # ── 长度路由 ──────────────────────────────────────────

    def _route_by_length(
        self,
        content: str,
        method: str,
        images: list[ExtractedImage] | None = None,
    ) -> ParseResult:
        # 清理无法在浏览器显示的本地图片引用，转为纯文字图注
        cleaned = clean_document_images(content)
        # 清理非法的 Unicode 代理字符，避免存储到 MySQL 时报错 3140
        cleaned = clean_surrogates(cleaned)
        truncated = cleaned[:MAX_CONTENT_CHARS]
        length = len(truncated)
        category = "short_doc" if length <= SHORT_DOC_THRESHOLD else "long_doc"
        return ParseResult(
            category=category,
            method=method,
            content=truncated,
            content_length=length,
            images=images or [],
        )


# ── 工具函数 ─────────────────────────────────────────────

def _is_valid_content(content: str) -> bool:
    text = (content or "").strip()
    if len(text) < 20:
        return False
    garbage = sum(1 for c in text if ord(c) > 0xFFFF)
    return garbage / max(len(text), 1) < 0.1


# ── 图片引用清理正则 ──────────────────────────────────────
# 匹配 ![alt](src) 中 src 为本地路径（不以 http/https/data: 开头）的图片
_LOCAL_IMG_RE = re.compile(
    r"!\[([^\]]*)\]\((?!https?://|data:)([^)]+)\)",
    re.IGNORECASE,
)
# 匹配 HTML <img src="local_path" ...>
_LOCAL_HTML_IMG_RE = re.compile(
    r'<img\b[^>]*\bsrc=["\'](?!https?://|data:)[^"\']*["\'][^>]*>',
    re.IGNORECASE,
)


def clean_document_images(content: str) -> str:
    """将文档中无法在浏览器访问的本地图片引用转为纯文本图注。

    - `![图4: 流程图](./images/fig4.png)` → `[图片: 图4: 流程图]`
    - `![](path/fig.png)` → `[图片（无图注）]`
    - HTML <img src="local"> → 直接移除
    保留 http/https/data: URL 的图片引用不变。
    """
    def _replace_md_img(m: re.Match) -> str:
        alt = (m.group(1) or "").strip()
        return f"[图片: {alt}]" if alt else "[图片（无图注）]"

    content = _LOCAL_IMG_RE.sub(_replace_md_img, content)
    content = _LOCAL_HTML_IMG_RE.sub("", content)
    return content


def replace_local_images_with_urls(content: str, image_urls: dict[str, str]) -> str:
    """将 markdown / HTML 中的本地图片引用替换为可访问 URL。"""
    if not content or not image_urls:
        return content

    def _normalize(path: str) -> list[str]:
        raw = (path or "").strip().strip('"\'')
        if not raw:
            return []
        raw = raw.replace("\\", "/")
        candidates = [raw]
        basename = os.path.basename(raw)
        if basename and basename not in candidates:
            candidates.append(basename)
        if raw.startswith("./"):
            stripped = raw[2:]
            if stripped not in candidates:
                candidates.append(stripped)
            stripped_base = os.path.basename(stripped)
            if stripped_base and stripped_base not in candidates:
                candidates.append(stripped_base)
        return candidates

    def _resolve(src: str) -> str | None:
        for key in _normalize(src):
            if key in image_urls:
                return image_urls[key]
        return None

    def _replace_md_img(m: re.Match) -> str:
        alt = m.group(1) or ""
        src = m.group(2) or ""
        url = _resolve(src)
        if not url:
            return m.group(0)
        return f"![{alt}]({url})"

    def _replace_html_img(m: re.Match) -> str:
        tag = m.group(0)
        src_m = re.search(r'\bsrc=["\']([^"\']+)["\']', tag, re.IGNORECASE)
        if not src_m:
            return tag
        url = _resolve(src_m.group(1))
        if not url:
            return tag
        return re.sub(
            r'(\bsrc=["\'])[^"\']+(["\'])',
            rf"\1{url}\2",
            tag,
            count=1,
            flags=re.IGNORECASE,
        )

    content = _LOCAL_IMG_RE.sub(_replace_md_img, content)
    content = _LOCAL_HTML_IMG_RE.sub(_replace_html_img, content)
    return content


# ── 图片压缩 ──────────────────────────────────────────────

MAX_IMAGE_BYTES = 2 * 1024 * 1024  # 2MB


def compress_image_if_needed(img: ExtractedImage) -> ExtractedImage:
    """如果图片超过 2MB，压缩为 JPEG quality=80"""
    if len(img.data) <= MAX_IMAGE_BYTES:
        return img
    try:
        from PIL import Image
        pil_img = Image.open(io.BytesIO(img.data))
        # 转换为 RGB（JPEG 不支持 RGBA）
        if pil_img.mode in ("RGBA", "P"):
            pil_img = pil_img.convert("RGB")
        buf = io.BytesIO()
        pil_img.save(buf, format="JPEG", quality=80)
        compressed = buf.getvalue()
        if len(compressed) < len(img.data):
            logger.info(
                "[ChatFileParser] 图片压缩: %s %dKB -> %dKB",
                img.name, len(img.data) // 1024, len(compressed) // 1024,
            )
            return ExtractedImage(
                data=compressed,
                name=img.name.rsplit(".", 1)[0] + ".jpg",
                caption=img.caption,
                content_type="image/jpeg",
            )
    except ImportError:
        logger.debug("[ChatFileParser] Pillow 未安装，跳过图片压缩")
    except Exception as e:
        logger.warning("[ChatFileParser] 图片压缩失败: %s", e)
    return img


# ── 图片 MIME 类型检测 ──────────────────────────────────

def _detect_image_content_type(data: bytes) -> str:
    """通过文件头检测图片 MIME 类型"""
    if data[:8] == b'\x89PNG\r\n\x1a\n':
        return "image/png"
    if data[:2] == b'\xff\xd8':
        return "image/jpeg"
    if data[:4] == b'GIF8':
        return "image/gif"
    if data[:2] == b'BM':
        return "image/bmp"
    if data[:4] == b'RIFF' and data[8:12] == b'WEBP':
        return "image/webp"
    return "image/png"  # 默认


def _guess_mime(ext: str) -> str:
    mapping = {
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".png": "image/png",
        ".gif": "image/gif",
        ".webp": "image/webp",
        ".bmp": "image/bmp",
        ".svg": "image/svg+xml",
    }
    return mapping.get(ext, "application/octet-stream")


def get_file_category_hint(filename: str) -> str:
    """快速判断文件大类（不解析内容）"""
    ext = Path(filename).suffix.lower()
    if ext in MULTIMODAL_EXTENSIONS:
        return "multimodal"
    if ext in PYTHON_PARSEABLE_EXTENSIONS | RICH_DOCUMENT_EXTENSIONS:
        return "document"
    return "unknown"


def clean_surrogates(val):
    """递归清理字典/列表/字符串中的 unpaired surrogates，避免 MySQL JSON column 插入报错 3140"""
    if isinstance(val, str):
        try:
            return val.encode("utf-16", errors="surrogatepass").decode("utf-16", errors="surrogatepass").encode("utf-8", errors="replace").decode("utf-8")
        except Exception:
            return val
    elif isinstance(val, dict):
        return {k: clean_surrogates(v) for k, v in val.items()}
    elif isinstance(val, list):
        return [clean_surrogates(x) for x in val]
    return val
