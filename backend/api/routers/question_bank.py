"""教师题库管理路由（基础力学课程智能体专属）。

仅教师/管理员可访问。批次是两段式：上传即解析（返回逐题校验报告），教师
预览修正后 confirm 才写题目表。题库为课程共享教学资产，所有教师可管理，
审计 + created_by 追溯，不做属主隔离。
"""

from __future__ import annotations

import asyncio
import io

from fastapi import APIRouter, Depends, File, HTTPException, Query, Request, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel, Field
from sqlalchemy.ext.asyncio import AsyncSession

from backend.api.dependencies.auth import ApiKeyAuth, require_teacher
from backend.application.question_bank import (
    QuestionBankError,
    append_question_image,
    batch_to_dict,
    confirm_batch,
    create_batch,
    create_knowledge_point,
    create_question,
    delete_batch,
    delete_knowledge_point,
    get_batch,
    get_question,
    list_knowledge_points,
    question_page,
    question_to_dict,
    remove_question_image,
    update_knowledge_point,
    update_question,
)
from backend.config.loggers import get_logger
from backend.db.connection import get_async_session
from backend.db.models import QuestionBankKnowledgePoint
from backend.services.admin_audit import write_admin_audit

logger = get_logger(__name__)

router = APIRouter(prefix="/teacher/question-bank", tags=["教师-题库管理"])

_MAX_UPLOAD_READ = 20 * 1024 * 1024 + 1
_MAX_IMAGE_BYTES = 10 * 1024 * 1024
_IMAGE_EXTS = {"png", "jpg", "jpeg", "gif", "webp"}


async def _audit(
    session: AsyncSession,
    request: Request,
    auth: ApiKeyAuth,
    *,
    action: str,
    target_type: str,
    target_id: str | None = None,
    detail: dict | None = None,
) -> None:
    try:
        await write_admin_audit(
            session,
            request,
            auth,
            action=action,
            target_type=target_type,
            target_id=target_id,
            detail=detail,
        )
        await session.commit()
    except Exception as exc:  # noqa: BLE001
        try:
            await session.rollback()
        except Exception:  # noqa: BLE001
            pass
        logger.warning("[QuestionBank] 审计写入失败（不影响响应）: %s", exc)


def _to_http(exc: QuestionBankError) -> HTTPException:
    return HTTPException(status_code=422, detail=str(exc))


class QuestionUpdateRequest(BaseModel):
    model_config = {"extra": "forbid"}

    question_type: str | None = Field(default=None, max_length=16)
    stem: str | None = Field(default=None, max_length=8000)
    options: list[dict[str, str]] | None = Field(default=None, max_length=10)
    answer: str | None = Field(default=None, max_length=20000)
    analysis: str | None = Field(default=None, max_length=20000)
    score: str | None = Field(default=None, max_length=32)
    difficulty: str | None = Field(default=None, max_length=16)
    chapter: str | None = Field(default=None, max_length=128)
    knowledge_points: list[str] | None = Field(default=None, max_length=20)
    source_name: str | None = Field(default=None, max_length=255)
    status: str | None = Field(default=None, max_length=16)


class QuestionCreateRequest(QuestionUpdateRequest):
    question_type: str = Field(max_length=16)
    stem: str = Field(max_length=8000)


class KnowledgePointCreateRequest(BaseModel):
    model_config = {"extra": "forbid"}

    name: str = Field(max_length=128)
    chapter: str | None = Field(default=None, max_length=128)


class KnowledgePointUpdateRequest(BaseModel):
    model_config = {"extra": "forbid"}

    name: str | None = Field(default=None, max_length=128)
    chapter: str | None = Field(default=None, max_length=128)
    sort: int | None = None


def _build_template_docx() -> bytes:
    """现场生成标准模板 docx（示例题 + 内嵌示例受力图），不把二进制放进 Git。"""
    from docx import Document
    from docx.shared import Inches

    from backend.application.question_bank.template_figure import (
        build_template_figure_png,
    )

    document = Document()
    document.add_heading("基础力学课程智能体 · 试题上传模板", level=1)
    document.add_paragraph(
        "填写说明：每题一个块，字段用固定标记开头；题与题之间用独立的 ==== 分隔行隔开；"
        "受力图/结构简图直接粘贴在【题干】内容位置（下方示例题已内嵌一张受力图，"
        "可参照其位置粘贴）；公式建议 LaTeX 文本或图片；答案与解析须教师自填，系统不生成。"
        "难度仅支持 简单/中等/较难。本说明段落与示例题不会被解析入库的模板结构冲突，"
        "可直接修改示例后上传。"
    )
    sample = [
        ("【题型】", "选择题"),
        ("【题干】", "图示简支梁跨度 l=4 m，跨中受集中力 P=10 kN（受力图见下），"
                     "其跨中截面的最大弯矩为（　）。"),
        ("【选项】", "A. PL/4　B. PL/8　C. PL/2　D. PL"),
        ("【答案】", "A"),
        ("【解析】", "跨中集中力作用下 Mmax=PL/4=10kN×4m/4=10kN·m，故选 A。"),
        ("【难度】", "简单"),
        ("【知识点】", "弯曲内力；弯矩图"),
        ("【章节】", "第 4 章 弯曲内力"),
        ("【来源】", "2024-2025 期末卷 第 3 题"),
    ]
    for marker, content in sample:
        document.add_paragraph(f"{marker}{content}")
        if marker == "【题干】":
            # 示例受力图：与真实上传同一路径——图在题干标记之后、选项之前，
            # 解析时锚定到本题（word/media 内嵌，非外链）。
            document.add_picture(
                io.BytesIO(build_template_figure_png()), width=Inches(3.9)
            )
    document.add_paragraph("====")
    document.add_paragraph("【题型】计算题")
    document.add_paragraph("【题干】悬臂梁自由端受集中力 F=2kN，长 L=3m，求固定端弯矩。")
    document.add_paragraph("【答案】M=F·L=6kN·m（上侧受拉/下侧受拉按坐标系约定）")
    document.add_paragraph("【解析】固定端弯矩等于荷载对固定端之矩 M=F×L=2×3=6kN·m。")
    document.add_paragraph("【难度】中等")
    document.add_paragraph("【知识点】悬臂梁；弯曲内力")
    document.add_paragraph("【章节】第 4 章 弯曲内力")
    document.add_paragraph("【来源】课后习题 4-2")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@router.get("/template", summary="下载试题上传模板 docx")
async def download_template(
    request: Request,
    session: AsyncSession = Depends(get_async_session),
    auth: ApiKeyAuth = Depends(require_teacher),
):
    # 首次渲染要导入 matplotlib 并画示例受力图（约秒级），放线程池执行，
    # 不阻塞事件循环上所有进行中的 SSE 流。
    content = await asyncio.to_thread(_build_template_docx)
    await _audit(
        session, request, auth,
        action="teacher.question_bank.template", target_type="question_bank",
    )
    return Response(
        content=content,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={
            "Content-Disposition": 'attachment; filename="question-bank-template.docx"'
        },
    )


@router.post("/batches", status_code=201, summary="上传 docx 并解析")
async def upload_batch(
    request: Request,
    file: UploadFile = File(...),
    session: AsyncSession = Depends(get_async_session),
    auth: ApiKeyAuth = Depends(require_teacher),
):
    data = await file.read(_MAX_UPLOAD_READ)
    if len(data) >= _MAX_UPLOAD_READ:
        raise HTTPException(status_code=413, detail="文件超过 20MB 上限")
    try:
        batch = await create_batch(
            session,
            file_name=file.filename or "upload.docx",
            file_bytes=data,
            created_by=auth.user_id,
        )
    except QuestionBankError as exc:
        raise _to_http(exc) from exc
    await _audit(
        session, request, auth,
        action="teacher.question_bank.batch.upload",
        target_type="question_bank_batch", target_id=batch.id,
        detail={"file_name": batch.file_name, "total": batch.total_count},
    )
    return batch_to_dict(batch, include_questions=True)


@router.get("/batches/{batch_id}", summary="查看批次解析结果")
async def get_batch_detail(
    batch_id: str,
    request: Request,
    session: AsyncSession = Depends(get_async_session),
    auth: ApiKeyAuth = Depends(require_teacher),
):
    batch = await get_batch(session, batch_id)
    if batch is None:
        raise HTTPException(status_code=404, detail="批次不存在")
    return batch_to_dict(batch, include_questions=True)


@router.post("/batches/{batch_id}/confirm", summary="确认批次入库")
async def confirm_batch_endpoint(
    batch_id: str,
    request: Request,
    session: AsyncSession = Depends(get_async_session),
    auth: ApiKeyAuth = Depends(require_teacher),
):
    batch = await get_batch(session, batch_id)
    if batch is None:
        raise HTTPException(status_code=404, detail="批次不存在")
    try:
        inserted, skipped_duplicates = await confirm_batch(session, batch)
    except QuestionBankError as exc:
        raise _to_http(exc) from exc
    await _audit(
        session, request, auth,
        action="teacher.question_bank.batch.confirm",
        target_type="question_bank_batch", target_id=batch.id,
        detail={"inserted": len(inserted), "skipped_duplicates": skipped_duplicates},
    )
    return {
        "success": True,
        "inserted": len(inserted),
        "skipped_duplicates": skipped_duplicates,
        "batch": batch_to_dict(batch, include_questions=True),
    }


@router.delete("/batches/{batch_id}", summary="废弃未确认批次")
async def delete_batch_endpoint(
    batch_id: str,
    request: Request,
    session: AsyncSession = Depends(get_async_session),
    auth: ApiKeyAuth = Depends(require_teacher),
):
    batch = await get_batch(session, batch_id)
    if batch is None:
        raise HTTPException(status_code=404, detail="批次不存在")
    try:
        await delete_batch(session, batch)
    except QuestionBankError as exc:
        raise _to_http(exc) from exc
    await _audit(
        session, request, auth,
        action="teacher.question_bank.batch.delete",
        target_type="question_bank_batch", target_id=batch_id,
    )
    return {"success": True}


@router.get("/questions", summary="题目列表（筛选/分页）")
async def list_questions(
    request: Request,
    question_type: str | None = Query(default=None),
    difficulty: str | None = Query(default=None),
    knowledge_point: str | None = Query(default=None, max_length=128),
    chapter: str | None = Query(default=None, max_length=128),
    status: str | None = Query(default=None),
    keyword: str | None = Query(default=None, max_length=200),
    limit: int = Query(20, ge=1, le=100),
    offset: int = Query(0, ge=0),
    session: AsyncSession = Depends(get_async_session),
    auth: ApiKeyAuth = Depends(require_teacher),
):
    rows, total = await question_page(
        session,
        question_type=question_type,
        difficulty=difficulty,
        knowledge_point=knowledge_point,
        chapter=chapter,
        status=status,
        keyword=keyword,
        limit=limit,
        offset=offset,
    )
    return {
        "items": [question_to_dict(row) for row in rows],
        "total": total,
        "limit": limit,
        "offset": offset,
    }


@router.post("/questions", status_code=201, summary="手工录入单题")
async def create_question_endpoint(
    payload: QuestionCreateRequest,
    request: Request,
    session: AsyncSession = Depends(get_async_session),
    auth: ApiKeyAuth = Depends(require_teacher),
):
    fields = payload.model_dump(exclude_unset=True)
    try:
        question = await create_question(
            session, created_by=auth.user_id, fields=fields
        )
    except QuestionBankError as exc:
        raise _to_http(exc) from exc
    await _audit(
        session, request, auth,
        action="teacher.question_bank.question.create",
        target_type="question_bank_question", target_id=question.id,
    )
    return question_to_dict(question)


@router.get("/questions/{question_id}", summary="题目详情（含答案解析）")
async def get_question_detail(
    question_id: str,
    request: Request,
    session: AsyncSession = Depends(get_async_session),
    auth: ApiKeyAuth = Depends(require_teacher),
):
    question = await get_question(session, question_id)
    if question is None:
        raise HTTPException(status_code=404, detail="题目不存在")
    return question_to_dict(question)


@router.patch("/questions/{question_id}", summary="编辑内容与标注")
async def update_question_endpoint(
    question_id: str,
    payload: QuestionUpdateRequest,
    request: Request,
    session: AsyncSession = Depends(get_async_session),
    auth: ApiKeyAuth = Depends(require_teacher),
):
    question = await get_question(session, question_id)
    if question is None:
        raise HTTPException(status_code=404, detail="题目不存在")
    changes = payload.model_dump(exclude_unset=True)
    try:
        question = await update_question(session, question, changes=changes)
    except QuestionBankError as exc:
        raise _to_http(exc) from exc
    await _audit(
        session, request, auth,
        action="teacher.question_bank.question.update",
        target_type="question_bank_question", target_id=question.id,
        detail={"fields": sorted(changes)},
    )
    return question_to_dict(question)


@router.delete("/questions/{question_id}", summary="下线题目（软删）")
async def delete_question_endpoint(
    question_id: str,
    request: Request,
    session: AsyncSession = Depends(get_async_session),
    auth: ApiKeyAuth = Depends(require_teacher),
):
    question = await get_question(session, question_id)
    if question is None:
        raise HTTPException(status_code=404, detail="题目不存在")
    try:
        question = await update_question(session, question, changes={"status": "disabled"})
    except QuestionBankError as exc:
        raise _to_http(exc) from exc
    await _audit(
        session, request, auth,
        action="teacher.question_bank.question.disable",
        target_type="question_bank_question", target_id=question_id,
    )
    return {"success": True}


@router.post("/questions/{question_id}/status", summary="发布/停用切换")
async def set_question_status(
    question_id: str,
    payload: QuestionUpdateRequest,
    request: Request,
    session: AsyncSession = Depends(get_async_session),
    auth: ApiKeyAuth = Depends(require_teacher),
):
    question = await get_question(session, question_id)
    if question is None:
        raise HTTPException(status_code=404, detail="题目不存在")
    status = (payload.status or "").strip()
    try:
        question = await update_question(session, question, changes={"status": status})
    except QuestionBankError as exc:
        raise _to_http(exc) from exc
    await _audit(
        session, request, auth,
        action="teacher.question_bank.question.status",
        target_type="question_bank_question", target_id=question_id,
        detail={"status": status},
    )
    return question_to_dict(question)


async def _read_image_upload(file: UploadFile) -> tuple[bytes, str]:
    data = await file.read(_MAX_IMAGE_BYTES + 1)
    if len(data) > _MAX_IMAGE_BYTES:
        raise HTTPException(status_code=413, detail="图片超过 10MB 上限")
    if not data:
        raise HTTPException(status_code=422, detail="图片为空")
    filename = str(file.filename or "image.png").lower()
    ext = filename.rsplit(".", 1)[-1] if "." in filename else ""
    if ext not in _IMAGE_EXTS:
        raise HTTPException(status_code=422, detail="仅支持 png/jpg/gif/webp 图片")
    if data[:8] == b"\x89PNG\r\n\x1a\n":
        ext = "png"
    elif data[:3] == b"\xff\xd8\xff":
        ext = "jpg"
    return data, ext


@router.post("/questions/{question_id}/images", status_code=201, summary="为题目补图")
async def add_question_image(
    question_id: str,
    request: Request,
    file: UploadFile = File(...),
    caption: str = Query(default="", max_length=200),
    session: AsyncSession = Depends(get_async_session),
    auth: ApiKeyAuth = Depends(require_teacher),
):
    question = await get_question(session, question_id)
    if question is None:
        raise HTTPException(status_code=404, detail="题目不存在")
    data, ext = await _read_image_upload(file)
    image = await append_question_image(
        session, question, blob=data, ext=ext, caption=caption
    )
    await _audit(
        session, request, auth,
        action="teacher.question_bank.question.image.add",
        target_type="question_bank_question", target_id=question_id,
        detail={"seq": image["seq"]},
    )
    return image


@router.delete("/questions/{question_id}/images/{seq}", summary="删除题目图片")
async def delete_question_image(
    question_id: str,
    seq: int,
    request: Request,
    session: AsyncSession = Depends(get_async_session),
    auth: ApiKeyAuth = Depends(require_teacher),
):
    question = await get_question(session, question_id)
    if question is None:
        raise HTTPException(status_code=404, detail="题目不存在")
    await remove_question_image(session, question, seq=seq)
    await _audit(
        session, request, auth,
        action="teacher.question_bank.question.image.delete",
        target_type="question_bank_question", target_id=question_id,
        detail={"seq": seq},
    )
    return {"success": True}


@router.get("/knowledge-points", summary="知识点字典列表")
async def list_kp_endpoint(
    request: Request,
    session: AsyncSession = Depends(get_async_session),
    auth: ApiKeyAuth = Depends(require_teacher),
):
    return {"items": await list_knowledge_points(session)}


@router.post("/knowledge-points", status_code=201, summary="新增知识点")
async def create_kp_endpoint(
    payload: KnowledgePointCreateRequest,
    request: Request,
    session: AsyncSession = Depends(get_async_session),
    auth: ApiKeyAuth = Depends(require_teacher),
):
    try:
        point = await create_knowledge_point(
            session, name=payload.name, chapter=payload.chapter, created_by=auth.user_id
        )
    except QuestionBankError as exc:
        raise _to_http(exc) from exc
    await _audit(
        session, request, auth,
        action="teacher.question_bank.kp.create",
        target_type="question_bank_kp", target_id=point.id,
        detail={"name": point.name},
    )
    return {
        "id": point.id, "name": point.name,
        "chapter": point.chapter, "parent_id": point.parent_id, "sort": point.sort,
    }


async def _get_kp(session: AsyncSession, kp_id: str) -> QuestionBankKnowledgePoint:
    point = await session.get(QuestionBankKnowledgePoint, kp_id)
    if point is None:
        raise HTTPException(status_code=404, detail="知识点不存在")
    return point


@router.patch("/knowledge-points/{kp_id}", summary="编辑知识点")
async def update_kp_endpoint(
    kp_id: str,
    payload: KnowledgePointUpdateRequest,
    request: Request,
    session: AsyncSession = Depends(get_async_session),
    auth: ApiKeyAuth = Depends(require_teacher),
):
    point = await _get_kp(session, kp_id)
    try:
        point = await update_knowledge_point(
            session, point, changes=payload.model_dump(exclude_unset=True)
        )
    except QuestionBankError as exc:
        raise _to_http(exc) from exc
    await _audit(
        session, request, auth,
        action="teacher.question_bank.kp.update",
        target_type="question_bank_kp", target_id=kp_id,
    )
    return {
        "id": point.id, "name": point.name,
        "chapter": point.chapter, "parent_id": point.parent_id, "sort": point.sort,
    }


@router.delete("/knowledge-points/{kp_id}", summary="删除知识点")
async def delete_kp_endpoint(
    kp_id: str,
    request: Request,
    session: AsyncSession = Depends(get_async_session),
    auth: ApiKeyAuth = Depends(require_teacher),
):
    point = await _get_kp(session, kp_id)
    try:
        await delete_knowledge_point(session, point)
    except QuestionBankError as exc:
        # 引用校验拒绝等业务性失败 → 422（此前是唯一未收敛的写端点，裸 500）。
        raise _to_http(exc) from exc
    await _audit(
        session, request, auth,
        action="teacher.question_bank.kp.delete",
        target_type="question_bank_kp", target_id=kp_id,
    )
    return {"success": True}
