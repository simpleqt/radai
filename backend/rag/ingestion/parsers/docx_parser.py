"""DOCX parser (附录 E.2) -- paragraphs, tables, embedded images -> typed blocks."""

from __future__ import annotations

import asyncio
import base64
from io import BytesIO
from typing import Any

from backend.rag.ingestion.ids import sha256_text, stable_block_id
from backend.rag.ingestion.models import IngestBlock, ParsedDocument, SourceDocument
from backend.rag.ingestion.parsers.base import BaseParser

_A_BLIP = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
_R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
_VML_IMAGEDATA = "{urn:schemas-microsoft-com:vml}imagedata"
_R_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"


class DocxParser(BaseParser):
    """Parse .docx into multi-type IngestBlock list.

    Phase-3 enhancements over Phase-1 inline docx logic:
    - Heading styles -> section_path hierarchy + searchable text blocks
    - Tables -> separate TableBlock (markdown format)
    - Embedded images -> ImageBlock in document order for downstream VLM
    """

    async def parse(
        self,
        content: bytes,
        source_document: SourceDocument,
    ) -> ParsedDocument:
        return await asyncio.to_thread(self._parse_sync, content, source_document)

    def _parse_sync(
        self,
        content: bytes,
        source_document: SourceDocument,
    ) -> ParsedDocument:
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError(
                "python-docx is required for DOCX parsing. pip install python-docx"
            ) from exc

        document = Document(BytesIO(content))
        blocks: list[IngestBlock] = []
        section_path: list[str] = []
        heading_levels: list[int] = []
        text_buffer: list[str] = []
        order = 0
        warnings: list[str] = []
        seen_image_rids: set[str] = set()

        def flush_text() -> None:
            nonlocal order
            body = "\n\n".join(text_buffer).strip()
            if not body:
                return
            blocks.append(_make_block(
                source_document, "text", body, list(section_path), order,
            ))
            order += 1

        def emit_images_from(element) -> None:
            nonlocal order
            for img_block in _images_from_element(
                document,
                element,
                source_document,
                section_path,
                order,
                warnings,
                seen_image_rids,
            ):
                blocks.append(img_block)
                order = img_block.order_index + 1

        paragraph_by_element = {
            id(paragraph._element): paragraph for paragraph in document.paragraphs
        }
        table_by_element = {id(table._tbl): table for table in document.tables}
        body_elements = list(document.element.body)

        for element in body_elements:
            tag = element.tag.split("}")[-1] if "}" in element.tag else element.tag

            if tag == "p":
                para = paragraph_by_element.get(id(element))
                if para is None:
                    emit_images_from(element)
                    continue
                text = (para.text or "").strip()
                style_name = (para.style.name or "") if para.style else ""

                level = _heading_level(style_name)
                if level:
                    flush_text()
                    text_buffer.clear()
                    while heading_levels and heading_levels[-1] >= level:
                        heading_levels.pop()
                        if section_path:
                            section_path.pop()
                    heading_levels.append(level)
                    title = text or f"Heading {level}"
                    section_path.append(title)
                    # Keep heading text searchable (aligned with MinerU content_list).
                    if text:
                        blocks.append(_make_block(
                            source_document, "text", text, list(section_path), order,
                        ))
                        order += 1
                    emit_images_from(element)
                    continue

                if text:
                    text_buffer.append(text)
                # Images may sit in a paragraph with or without visible text.
                if _element_image_rids(element):
                    flush_text()
                    text_buffer.clear()
                    emit_images_from(element)

            elif tag == "tbl":
                flush_text()
                text_buffer.clear()
                table = table_by_element.get(id(element))
                if table is not None:
                    md_table = _table_to_markdown(table)
                    if md_table:
                        blocks.append(_make_block(
                            source_document, "table", md_table, list(section_path), order,
                        ))
                        order += 1
                emit_images_from(element)

        flush_text()
        text_buffer.clear()

        # Orphan images not referenced from body (rare); append to preserve content.
        orphan_blocks = _extract_remaining_images(
            document, source_document, section_path, order, warnings, seen_image_rids,
        )
        blocks.extend(orphan_blocks)

        raw_text = "\n\n".join(b.content for b in blocks if b.block_type == "text")
        return ParsedDocument(
            source_document=source_document,
            blocks=blocks,
            raw_text=raw_text,
            parser_artifacts={"parser": "docx_v2"},
            warnings=warnings,
        )


def _heading_level(style_name: str) -> int | None:
    s = style_name.lower().strip()
    if s.startswith("heading"):
        rest = s.replace("heading", "").strip()
        if rest.isdigit():
            return int(rest)
    # Chinese Word styles: "标题 1" / "标题1"
    for cn in ("标题",):
        if s.startswith(cn):
            rest = s.replace(cn, "").strip()
            if rest.isdigit():
                return int(rest)
    return None


def _table_to_markdown(table) -> str:
    rows: list[list[str]] = []
    for row in table.rows:
        cells = [" ".join((cell.text or "").split()) for cell in row.cells]
        if any(cells):
            rows.append(cells)
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    normalized = [row + [""] * (width - len(row)) for row in rows]
    header = normalized[0]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(["---"] * width) + " |",
    ]
    for row in normalized[1:]:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)


def _element_image_rids(element) -> list[str]:
    rids: list[str] = []
    for blip in element.findall(f".//{_A_BLIP}"):
        rid = blip.get(_R_EMBED)
        if rid:
            rids.append(rid)
    for imagedata in element.findall(f".//{_VML_IMAGEDATA}"):
        rid = imagedata.get(_R_ID) or imagedata.get("r:id")
        if rid:
            rids.append(rid)
    return rids


def _image_block_from_rel(
    rel,
    sd: SourceDocument,
    section_path: list[str],
    order: int,
    warnings: list[str],
) -> IngestBlock | None:
    try:
        img_data = rel.target_part.blob
        img_b64 = base64.b64encode(img_data).decode("ascii")
        content_type = getattr(rel.target_part, "content_type", "image/png")
        ext = content_type.split("/")[-1] if "/" in content_type else "png"
        part_name = getattr(rel.target_part, "partname", f"image_{order}")
        content_text = f"[嵌入图片] {part_name}"
        return _make_block(
            sd, "image", content_text, list(section_path), order,
            metadata={
                "img_base64": img_b64,
                "img_base64_preview": img_b64[:200],
                "img_size_bytes": len(img_data),
                "content_type": content_type,
                "part_name": str(part_name),
                "ext": ext,
            },
        )
    except Exception as exc:
        warnings.append(f"图片提取失败: {exc}")
        return None


def _images_from_element(
    document,
    element,
    sd: SourceDocument,
    section_path: list[str],
    start_order: int,
    warnings: list[str],
    seen_image_rids: set[str],
) -> list[IngestBlock]:
    blocks: list[IngestBlock] = []
    order = start_order
    rels = document.part.rels
    for rid in _element_image_rids(element):
        if rid in seen_image_rids:
            continue
        rel = rels.get(rid)
        if rel is None:
            continue
        if "image" not in (rel.reltype or "").lower():
            continue
        seen_image_rids.add(rid)
        block = _image_block_from_rel(rel, sd, section_path, order, warnings)
        if block is not None:
            blocks.append(block)
            order += 1
    return blocks


def _extract_remaining_images(
    document,
    sd: SourceDocument,
    section_path: list[str],
    start_order: int,
    warnings: list[str],
    seen_image_rids: set[str],
) -> list[IngestBlock]:
    """Fallback for image parts never referenced from body traversal."""
    blocks: list[IngestBlock] = []
    order = start_order
    try:
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        for rid, rel in document.part.rels.items():
            if rid in seen_image_rids:
                continue
            if "image" not in (rel.reltype or "").lower() and not (
                hasattr(RT, "IMAGE") and rel.reltype == RT.IMAGE
            ):
                continue
            seen_image_rids.add(rid)
            block = _image_block_from_rel(rel, sd, section_path, order, warnings)
            if block is not None:
                blocks.append(block)
                order += 1
    except Exception as exc:
        warnings.append(f"DOCX 残留图片提取失败: {exc}")
    return blocks


def _make_block(
    sd: SourceDocument,
    block_type: str,
    content: str,
    section_path: list[str],
    order_index: int,
    metadata: dict[str, Any] | None = None,
) -> IngestBlock:
    block_id = stable_block_id(
        sd.doc_id, sd.commit_version, block_type,
        section_path=section_path,
        order_index=order_index,
    )
    return IngestBlock(
        block_id=block_id,
        doc_id=sd.doc_id,
        kb_id=sd.kb_id,
        kb_code=sd.kb_code,
        commit_version=sd.commit_version,
        block_type=block_type,
        content=content,
        section_path=list(section_path),
        order_index=order_index,
        source_hash=sd.source_hash,
        original_file_name=sd.original_file_name,
        object_storage_uri=sd.object_storage_uri,
        content_hash=sha256_text(content),
        stable_block_fingerprint=sha256_text(f"{block_type}:{content}"),
        metadata=metadata or {},
        job_id=sd.job_id,
    )
