"""Self-ingestion pipeline entry points."""

from __future__ import annotations

import asyncio
import os
import uuid
from datetime import datetime
from io import BytesIO
from pathlib import Path

from redis.exceptions import RedisError
from sqlalchemy import select, update

from backend.config import settings
from backend.config.loggers import get_logger
from backend.db.connection import async_session_context, retry_async_db_operation
from backend.db.models import (
    KnowledgeBase,
    KnowledgeBaseDocument,
    PendingFile,
    SelfIngestChunk,
    SelfIngestJob,
)
from backend.domain.knowledge_base.backend_policy import (
    SELF_BACKEND,
    should_use_self_ingestion,
)
from backend.rag.ingestion import states
from backend.rag.ingestion.block_builder import (
    build_chunk_quality_audit,
    build_chunks_from_blocks,
)
from backend.rag.ingestion.chunker import build_chunks
from backend.rag.ingestion.commit_writer import (
    CommitWriter,
    entity_embed_texts,
    media_chunks_from_chunks,
    visible_entities_from_extraction,
)
from backend.rag.ingestion.concurrency import commit_mysql_concurrency, semaphore
from backend.rag.ingestion.embedded_image_loader import load_and_enrich_embedded_images
from backend.rag.ingestion.embedding import (
    embed_media_texts,
    embed_texts,
    require_embedding_success,
)
from backend.rag.ingestion.extraction.llm_extractor import (
    SelfExtractionError,
    extract_chunks_schema_aware,
)
from backend.rag.ingestion.extraction.vision_extractor import ImageAsset
from backend.rag.ingestion.fencing import IngestFenceLostError
from backend.rag.ingestion.ids import stable_doc_id
from backend.rag.ingestion.kb_stats import refresh_knowledge_base_stats
from backend.rag.ingestion.models import SourceDocument
from backend.rag.ingestion.multimodal_summarizer import enrich_table_formula_blocks
from backend.rag.ingestion.parsers.parser_router import ParserRouter
from backend.rag.ingestion.staging_manager import (
    StagingPayload,
    delete_staging_payload,
    save_staging_payload,
)
from backend.rag.ingestion.storage.milvus_writer import MilvusChunkWriter
from backend.storage import download_minio_object_to_path, sha256_file

logger = get_logger(__name__)

TEXT_EXTENSIONS = {".md", ".markdown", ".txt", ".text"}
TEXT_MIME_TYPES = {"text/plain", "text/markdown", "text/x-markdown"}
PDF_EXTENSIONS = {".pdf"}
PDF_MIME_TYPES = {"application/pdf"}
DOCX_EXTENSIONS = {".docx"}
DOCX_MIME_TYPES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tiff"}
IMAGE_MIME_TYPES = {"image/png", "image/jpeg", "image/gif", "image/webp", "image/bmp", "image/tiff"}

STAGE_SKIP = "skipped"


def _core_mysql_retries() -> int:
    try:
        return max(int(getattr(settings, "self_ingestion_commit_mysql_retries", 8) or 8), 1)
    except (TypeError, ValueError):
        return 8


def _core_mysql_retry_base_delay() -> float:
    try:
        return max(float(getattr(settings, "self_ingestion_commit_mysql_retry_base_delay", 1.0) or 1.0), 0.1)
    except (TypeError, ValueError):
        return 1.0


class SelfIngestionNotReadyError(RuntimeError):
    """Raised when a self-ingestion task is routed before the writer is complete."""


class SelfIngestionPipeline:
    """Phase-1 self-ingestion pipeline for text-like documents."""

    async def ingest_pending_file(self, file_id: str, *, fence_check=None) -> bool:
        """两阶段可见性提交 + 细粒度阶段流转 + commit_log 审计。

        S1: MySQL 不再直接写 visible，失败回滚为 failed。
        S2: 旧版本在新版本 mark_visible 之前保持原可见状态，消除查询空窗。
        S4: 文档注册在事务内 SELECT FOR UPDATE，消除并发去重竞态。
        A: 外部索引（Milvus/Redis/Nebula）部分失败 → JOB_PARTIAL_SUCCESS，repair_worker 补偿。
        C: job.stage 在 PARSING/CHUNKING/EXTRACTING/EMBEDDING/COMMITTING_*/MARKING_VISIBLE 间流转。
        D: commit_log 记录各阶段状态与耗时，写入 SelfIngestJob.commit_log_json 与 Redis。
        """
        job_id = str(uuid.uuid4())
        doc_id: str | None = None
        kb_code: str | None = None
        vector_count = 0
        commit_version = 1
        commit_log: list[dict] = []
        metrics: dict = {}
        downloaded_path: Path | None = None

        async def _assert_fence() -> None:
            if fence_check is not None:
                await fence_check()

        def _log_stage(stage: str, status: str, *, detail: str | None = None, **kw) -> None:
            entry: dict = {"stage": stage, "status": status, "at": datetime.now().isoformat()}
            if detail:
                entry["detail"] = detail[:300]
            entry.update(kw)
            commit_log.append(entry)

        async def _set_stage(stage: str) -> None:
            await _assert_fence()
            async with async_session_context() as session:
                await session.execute(
                    update(SelfIngestJob)
                    .where(SelfIngestJob.job_id == job_id)
                    .values(stage=stage, updated_at=datetime.now())
                )
                await session.commit()

        try:
            await _assert_fence()
            # 1. 加载 pending + kb，预写 job 行
            async with async_session_context() as session:
                row = (
                    await session.execute(
                        select(PendingFile, KnowledgeBase)
                        .join(KnowledgeBase, PendingFile.kb_id == KnowledgeBase.kb_id)
                        .where(PendingFile.file_id == file_id)
                    )
                ).one_or_none()
                if not row:
                    logger.warning("[self-ingestion] pending file not found: %s", file_id)
                    return False
                pending_file, kb = row
                if not should_use_self_ingestion(settings, kb=kb, pending_file=pending_file):
                    raise RuntimeError(f"self_ingestion 未启用或未授权: kb_id={pending_file.kb_id}")
                if pending_file.status != 1:
                    logger.info(
                        "[self-ingestion] skip non-processing file: file_id=%s status=%s",
                        file_id, pending_file.status,
                    )
                    return False
                await self._start_job(session, job_id, pending_file)
                await session.commit()

            kb_code = kb.kb_code
            download_dir = (
                Path(__file__).resolve().parents[2]
                / "data"
                / "cache"
                / "self_ingestion_downloads"
            )
            downloaded_path = download_dir / (
                f"{file_id}_{Path(pending_file.file_name).name}"
            )
            await asyncio.to_thread(
                download_minio_object_to_path,
                pending_file.minio_path,
                downloaded_path,
                bucket_name=pending_file.object_bucket,
            )
            file_size = downloaded_path.stat().st_size
            source_hash = pending_file.source_hash or await asyncio.to_thread(
                sha256_file,
                downloaded_path,
            )
            file_ext = (
                pending_file.file_ext
                or Path(pending_file.file_name or "").suffix
            ).lower()
            mime_type = (
                pending_file.mime_type or pending_file.file_type or ""
            ).lower()
            is_pdf = file_ext in PDF_EXTENSIONS or mime_type in PDF_MIME_TYPES
            try:
                max_in_memory = max(
                    1024 * 1024,
                    int(
                        os.getenv(
                            "SELF_INGESTION_MAX_IN_MEMORY_FILE_BYTES",
                            "134217728",
                        )
                    ),
                )
            except (TypeError, ValueError):
                max_in_memory = 128 * 1024 * 1024
            if is_pdf:
                content: bytes | Path = downloaded_path
            elif file_size <= max_in_memory:
                content = await asyncio.to_thread(downloaded_path.read_bytes)
            else:
                raise RuntimeError(
                    f"文件 {file_size} bytes 超过非 PDF 内存解析上限 "
                    f"{max_in_memory} bytes"
                )
            logical_path = "/".join(
                part.strip("/")
                for part in (
                    pending_file.category_path or "",
                    pending_file.file_name or "",
                )
                if part
            )
            doc_id = stable_doc_id(
                pending_file.kb_id,
                source_hash,
                pending_file.file_name,
                logical_path=logical_path,
            )

            # 2. 注册文档（事务内 SELECT FOR UPDATE，S4）
            async def _register_doc_transaction() -> tuple[str, int]:
                async with async_session_context() as session:
                    result = await self._register_doc(
                        session,
                        pending_file=pending_file,
                        job_id=job_id,
                        doc_id=doc_id,
                        source_hash=source_hash,
                    )
                    await session.commit()
                    return result

            commit_sem = await semaphore(
                "self_ingestion_commit_mysql",
                commit_mysql_concurrency(),
            )
            async with commit_sem:
                await _assert_fence()
                action, commit_version = await retry_async_db_operation(
                    _register_doc_transaction,
                    attempts=_core_mysql_retries(),
                    base_delay=_core_mysql_retry_base_delay(),
                    label=f"self-ingestion register_doc file_id={file_id}",
                    logger=logger,
                )
                await _assert_fence()

            if action == "skip_existing":
                logger.info(
                    "[self-ingestion] 跳过重复文件: kb=%s file=%s doc=%s hash=%s",
                    kb.kb_code, pending_file.file_name, doc_id, source_hash[:16],
                )
                _log_stage(STAGE_SKIP, "skipped", detail="相同内容已入库")
                async with async_session_context() as session:
                    await session.execute(
                        update(PendingFile)
                        .where(PendingFile.file_id == file_id)
                        .values(
                            status=2, source_hash=source_hash, ingestion_backend=SELF_BACKEND,
                            error_message="skip_existing: 相同内容已入库",
                            last_ingest_finished_at=datetime.now(),
                        )
                    )
                    await self._finalize_job(
                        session, job_id, status=states.JOB_COMPLETED, stage=STAGE_SKIP,
                        doc_id=doc_id, source_hash=source_hash,
                        commit_log=commit_log, metrics=metrics,
                    )
                    await session.commit()
                return True

            # 3. PARSING (multi-modal parser router, Phase-3)
            await _set_stage(states.STAGE_PARSING)
            t0 = datetime.now()
            sd = SourceDocument(
                doc_id=doc_id,
                kb_id=pending_file.kb_id,
                kb_code=kb.kb_code,
                file_id=file_id,
                job_id=job_id,
                subject=kb.subject,
                original_file_name=pending_file.file_name,
                file_ext=(pending_file.file_ext or Path(pending_file.file_name or "").suffix).lower(),
                mime_type=(pending_file.mime_type or pending_file.file_type or ""),
                file_size=file_size,
                source_hash=source_hash,
                object_storage_uri=pending_file.original_uri or pending_file.minio_path or "",
                parser_output_uri=(
                    pending_file.minio_path
                    if (pending_file.minio_path or "").lower().endswith("_content_list.json")
                    else None
                ),
                commit_version=commit_version,
            )
            embedded_assets: dict[str, ImageAsset] = {}
            text = ""
            is_pdf = sd.file_ext in PDF_EXTENSIONS or sd.mime_type in PDF_MIME_TYPES
            try:
                router = ParserRouter()
                parsed = await router.parse(content, sd)
                block_types = {b.block_type for b in parsed.blocks}
                metrics["parse_ms"] = int((datetime.now() - t0).total_seconds() * 1000)
                metrics["block_count"] = len(parsed.blocks)
                metrics["block_types"] = sorted(block_types)
                _log_stage(states.STAGE_PARSING, "ok",
                           blocks=len(parsed.blocks), types=sorted(block_types),
                           warnings=parsed.warnings[:5] if parsed.warnings else [])
                if parsed.blocks:
                    embedded_assets, embedded_stats = await load_and_enrich_embedded_images(
                        parsed.blocks,
                        source_document=sd,
                        parser_artifacts=parsed.parser_artifacts,
                    )
                    metrics["embedded_images"] = embedded_stats
                else:
                    # Successful parse with zero blocks (e.g. empty/odd DOCX).
                    # Fall back to text extraction for non-PDF; never leave `text` unset.
                    if is_pdf:
                        raise RuntimeError(
                            f"PDF 解析结果为空: {pending_file.file_name}"
                        )
                    logger.warning(
                        "[self-ingestion] 解析结果为空，回退纯文本: file=%s",
                        pending_file.file_name,
                    )
                    assert isinstance(content, bytes)
                    text = await self._extract_text(pending_file, content, file_id)
                    parsed = None
                    _log_stage(
                        states.STAGE_PARSING,
                        "fallback_text_empty_blocks",
                        chars=len(text or ""),
                    )
            except Exception as parse_exc:
                if is_pdf:
                    metrics["parse_ms"] = int((datetime.now() - t0).total_seconds() * 1000)
                    _log_stage(
                        states.STAGE_PARSING,
                        "mineru_failed",
                        detail=str(parse_exc)[:300],
                    )
                    raise
                logger.warning(
                    "[self-ingestion] 多模态解析失败，回退纯文本: file=%s error=%s",
                    pending_file.file_name, parse_exc,
                )
                assert isinstance(content, bytes)
                text = await self._extract_text(pending_file, content, file_id)
                parsed = None
                metrics["parse_ms"] = int((datetime.now() - t0).total_seconds() * 1000)
                _log_stage(states.STAGE_PARSING, "fallback_text", chars=len(text or ""))

            if parsed is not None and parsed.blocks:
                t_summary = datetime.now()
                try:
                    multimodal_summary = await enrich_table_formula_blocks(
                        parsed.blocks,
                        subject=kb.subject,
                        require_llm=False,
                    )
                    multimodal_summary["duration_ms"] = int(
                        (datetime.now() - t_summary).total_seconds() * 1000
                    )
                    metrics["multimodal_summary"] = multimodal_summary
                    if multimodal_summary.get("target_count"):
                        status = "ok" if not multimodal_summary.get("failed_count") else "partial"
                        _log_stage(
                            "multimodal_summary",
                            status,
                            table_summarized=multimodal_summary.get("table_summarized", 0),
                            formula_summarized=multimodal_summary.get("formula_summarized", 0),
                            failed=multimodal_summary.get("failed_count", 0),
                        )
                except Exception as summary_exc:
                    logger.warning(
                        "[self-ingestion] 表格/公式摘要增强失败，继续入库: file=%s error=%s",
                        pending_file.file_name,
                        summary_exc,
                    )
                    metrics["multimodal_summary"] = {
                        "enabled": False,
                        "error": str(summary_exc)[:300],
                    }

            # 4. CHUNKING (multi-modal block builder, Phase-3)
            await _set_stage(states.STAGE_CHUNKING)
            t0 = datetime.now()
            chunking_config = settings.rag_ingestion_config.get("chunking", {})
            chunk_size = int(chunking_config.get("rag_chunk_token_size") or 800)
            chunk_overlap = int(chunking_config.get("rag_chunk_overlap_token_size") or 120)
            if parsed is not None and parsed.blocks:
                chunks = build_chunks_from_blocks(
                    parsed,
                    file_name=pending_file.file_name,
                    file_path=pending_file.original_uri or pending_file.minio_path,
                    source_hash=source_hash,
                    document_date=pending_file.document_date,
                    category_path=pending_file.category_path,
                    commit_version=commit_version,
                    chunk_size_tokens=chunk_size,
                    chunk_overlap_tokens=chunk_overlap,
                )
            else:
                chunks = build_chunks(
                    text,
                    doc_id=doc_id,
                    file_name=pending_file.file_name,
                    file_path=pending_file.original_uri or pending_file.minio_path,
                    source_hash=source_hash,
                    document_date=pending_file.document_date,
                    category_path=pending_file.category_path,
                    commit_version=commit_version,
                    chunk_size_tokens=chunk_size,
                    chunk_overlap_tokens=chunk_overlap,
                )
            metrics["chunk_ms"] = int((datetime.now() - t0).total_seconds() * 1000)
            metrics["chunk_count"] = len(chunks)
            chunk_type_counts = {}
            for c in chunks:
                chunk_type_counts[c.block_type] = chunk_type_counts.get(c.block_type, 0) + 1
            metrics["chunk_type_counts"] = chunk_type_counts
            metrics["chunk_quality"] = build_chunk_quality_audit(
                chunks,
                blocks=list(parsed.blocks) if parsed is not None and parsed.blocks else None,
                chunk_size_tokens=chunk_size,
            )
            if not chunks:
                raise RuntimeError("解析后没有可入库 chunk")
            _log_stage(
                states.STAGE_CHUNKING,
                "ok",
                chunks=len(chunks),
                types=chunk_type_counts,
                short_chunks=metrics["chunk_quality"].get("char_len_le_30", 0),
                filtered_text_blocks=metrics["chunk_quality"].get("text_block_filtered_count", 0),
            )

            # 5. EXTRACTING（文本 LLM + 图片 VLM，强制模型抽取）
            await _set_stage(states.STAGE_EXTRACTING)
            t0 = datetime.now()
            image_assets = self._build_image_assets(
                parsed,
                content=content if isinstance(content, bytes) else b"",
                file_name=pending_file.file_name or "",
                file_ext=(pending_file.file_ext or Path(pending_file.file_name or "").suffix).lower(),
                mime_type=(pending_file.mime_type or pending_file.file_type or ""),
                embedded_assets=embedded_assets,
            )
            try:
                extraction, extraction_audit = await extract_chunks_schema_aware(
                    kb.kb_id,
                    chunks,
                    kb_code=kb.kb_code,
                    subject=kb.subject,
                    blocks=list(parsed.blocks) if parsed is not None and parsed.blocks else None,
                    image_assets=image_assets,
                    require_llm=True,
                )
            except SelfExtractionError as extract_exc:
                raise RuntimeError(f"实体/关系模型抽取失败: {extract_exc}") from extract_exc
            metrics["extract_ms"] = int((datetime.now() - t0).total_seconds() * 1000)
            metrics["entity_count"] = len(extraction.entities)
            metrics["relation_count"] = len(extraction.relations)
            metrics["extraction"] = extraction_audit
            _log_stage(states.STAGE_EXTRACTING, "ok",
                       entities=len(extraction.entities), relations=len(extraction.relations),
                       schema=extraction_audit.get("schema_version"),
                       extractor=extraction_audit.get("extractor"))

            # 6. EMBEDDING
            await _set_stage(states.STAGE_EMBEDDING)
            t0 = datetime.now()
            visible_entities = visible_entities_from_extraction(extraction)
            embed_entities_during_ingest = bool(
                getattr(settings, "self_ingestion_embed_entities_during_ingest", True)
            )
            embedding_required = require_embedding_success()
            entity_vectors: list[list[float]] = []
            media_vectors: list[list[float]] = []
            media_chunks = media_chunks_from_chunks(chunks)
            media_task: asyncio.Task[list[list[float]]] | None = None

            async def _embed_media_safe() -> list[list[float]]:
                if not media_chunks:
                    return []
                try:
                    return await embed_media_texts([chunk.content for chunk in media_chunks])
                except Exception as media_exc:
                    logger.warning(
                        "[self-ingestion] media embedding 失败，跳过图向量: file_id=%s media_chunks=%d error=%s",
                        file_id,
                        len(media_chunks),
                        media_exc,
                    )
                    return []

            try:
                if media_chunks:
                    media_task = asyncio.create_task(_embed_media_safe())
                if embed_entities_during_ingest and visible_entities:
                    vectors, entity_vectors = await asyncio.gather(
                        embed_texts([chunk.content for chunk in chunks]),
                        embed_texts(entity_embed_texts(visible_entities)),
                    )
                else:
                    vectors = await embed_texts([chunk.content for chunk in chunks])
            except Exception as exc:
                if media_task is not None and not media_task.done():
                    media_task.cancel()
                if embedding_required:
                    raise RuntimeError(f"embedding 失败，已按配置终止入库: {exc}") from exc
                logger.warning(
                    "[self-ingestion] embedding 失败，降级为 MySQL-only 检索: file_id=%s error=%s",
                    file_id, exc,
                )
                vectors = []
                entity_vectors = []
                media_vectors = []
                _log_stage(states.STAGE_EMBEDDING, "degraded", detail=str(exc)[:200])
            else:
                if len(vectors) != len(chunks):
                    if embedding_required:
                        raise RuntimeError(
                            f"embedding 数量不齐: chunks={len(chunks)} vectors={len(vectors)}"
                        )
                    logger.warning(
                        "[self-ingestion] embedding 数量不齐 chunks=%d vectors=%d",
                        len(chunks), len(vectors),
                    )
                    _log_stage(
                        states.STAGE_EMBEDDING,
                        "degraded",
                        detail=f"vectors={len(vectors)} chunks={len(chunks)}",
                    )
                    vectors = []
                    entity_vectors = []
                    media_vectors = []
                elif entity_vectors and len(entity_vectors) != len(visible_entities):
                    logger.warning(
                        "[self-ingestion] entity embedding 数量不齐 entities=%d vectors=%d",
                        len(visible_entities),
                        len(entity_vectors),
                    )
                    entity_vectors = []
                    _log_stage(
                        states.STAGE_EMBEDDING,
                        "ok",
                        vectors=len(vectors),
                        entity_vectors=0,
                    )
                else:
                    if media_task is not None:
                        media_vectors = await media_task
                        if media_vectors and len(media_vectors) != len(media_chunks):
                            logger.warning(
                                "[self-ingestion] media embedding 数量不齐 media_chunks=%d vectors=%d",
                                len(media_chunks),
                                len(media_vectors),
                            )
                            media_vectors = []
                    _log_stage(
                        states.STAGE_EMBEDDING,
                        "ok",
                        vectors=len(vectors),
                        entity_vectors=len(entity_vectors),
                        media_vectors=len(media_vectors),
                    )
            metrics["embed_ms"] = int((datetime.now() - t0).total_seconds() * 1000)

            # 7. STAGING：构建 StagingPayload 并持久化到 MinIO（§13）
            await _set_stage(states.STAGE_STAGING)
            t0 = datetime.now()
            payload = StagingPayload(
                job_id=job_id,
                kb_id=pending_file.kb_id,
                kb_code=kb.kb_code,
                doc_id=doc_id,
                commit_version=commit_version,
                source_document=sd,
                blocks=list(parsed.blocks) if parsed is not None and parsed.blocks else [],
                chunks=chunks,
                vectors=vectors,
                entity_vectors=entity_vectors,
                media_vectors=media_vectors,
                extraction=extraction,
                extraction_audit=extraction_audit,
                quality_audit=dict(metrics.get("chunk_quality") or {}),
                candidates=list(extraction_audit.get("candidate_facts") or []),
                dropped_logs=list(extraction_audit.get("dropped_facts") or []),
                metrics=dict(metrics),
            )
            try:
                await _assert_fence()
                payload_uri = await save_staging_payload(payload)
                await _assert_fence()
                async with async_session_context() as session:
                    await session.execute(
                        update(SelfIngestJob)
                        .where(SelfIngestJob.job_id == job_id)
                        .values(staging_uri=payload_uri, updated_at=datetime.now())
                    )
                    await session.commit()
                _log_stage(states.STAGE_STAGING, "ok", uri=payload_uri)
            except Exception as staging_exc:
                logger.warning(
                    "[self-ingestion] staging 持久化失败，继续 inline 提交: %s", staging_exc,
                )
                payload_uri = None
                _log_stage(states.STAGE_STAGING, "skipped", detail=str(staging_exc)[:200])
            metrics["staging_ms"] = int((datetime.now() - t0).total_seconds() * 1000)

            # 8-11. COMMIT：通过 CommitWriter 状态机写入四库（§14）
            commit_t0 = datetime.now()
            writer = CommitWriter()
            await _assert_fence()
            commit_result = await writer.commit(
                payload, kb=kb, pending_file=pending_file,
                set_stage_fn=_set_stage,
                log_stage_fn=_log_stage,
                fence_check=fence_check,
            )
            await _assert_fence()
            vector_count = commit_result.vector_count
            metrics["commit_ms"] = int((datetime.now() - commit_t0).total_seconds() * 1000)
            metrics["entity_vector_count"] = commit_result.entity_vector_count
            metrics["external_failures"] = commit_result.external_failures
            commit_log.extend(commit_result.stage_log)

            final_status = commit_result.final_status
            final_stage = commit_result.final_stage
            metrics["total_ms"] = sum(
                v for k, v in metrics.items() if isinstance(v, (int, float)) and k.endswith("_ms") and k != "total_ms"
            )

            if not commit_result.success:
                raise RuntimeError(
                    f"CommitWriter 核心阶段失败: {commit_result.error}"
                )

            async with async_session_context() as session:
                await self._finalize_job(
                    session, job_id, status=final_status, stage=final_stage,
                    doc_id=doc_id, source_hash=source_hash,
                    commit_log=commit_log, metrics=metrics,
                    extra={
                        "chunk_count": len(chunks),
                        "vector_count": vector_count,
                        "entity_count": len(extraction.entities),
                        "relation_count": len(extraction.relations),
                        "commit_version": commit_version,
                        "staging_uri": payload_uri,
                    },
                )
                kb_row = (
                    await session.execute(
                        select(KnowledgeBase).where(KnowledgeBase.kb_id == pending_file.kb_id)
                    )
                ).scalar_one()
                await self._refresh_kb_stats(session, kb_row)
                await session.commit()
            await self._write_commit_log_redis(kb.kb_code, job_id, commit_log, final_status)

            if commit_result.external_failures:
                logger.warning(
                    "[self-ingestion] 部分成功: kb=%s doc=%s ver=%d failures=%s staging=%s",
                    kb.kb_code, doc_id, commit_version, commit_result.external_failures, payload_uri,
                )
            else:
                if payload_uri and await delete_staging_payload(payload_uri):
                    async with async_session_context() as session:
                        await session.execute(
                            update(SelfIngestJob)
                            .where(SelfIngestJob.job_id == job_id)
                            .values(staging_uri=None, updated_at=datetime.now())
                        )
                        await session.commit()
                    payload_uri = None
                logger.info(
                    "[self-ingestion] 入库成功: kb=%s file=%s doc=%s ver=%d chunks=%d vectors=%d",
                    kb.kb_code, pending_file.file_name, doc_id, commit_version,
                    len(chunks), vector_count,
                )
            try:
                from backend.rag.ingestion.metrics import record_job_status
                record_job_status(kb.kb_code, final_status)
            except Exception:
                pass
            return True
        except IngestFenceLostError:
            logger.warning(
                "[self-ingestion] fencing token 已失效，旧 worker 停止且不写回终态: file_id=%s",
                file_id,
            )
            raise
        except asyncio.CancelledError:
            interrupted = RuntimeError("入库 Worker 停止，任务已安全重新排队")
            logger.warning("[self-ingestion] 入库被中断并重新排队: file_id=%s", file_id)
            _log_stage(states.STAGE_FAILED, "interrupted", detail=str(interrupted))
            await asyncio.shield(
                self._rollback_failure(
                    file_id=file_id,
                    job_id=job_id,
                    doc_id=doc_id,
                    kb_code=kb_code,
                    commit_version=commit_version,
                    vector_count=vector_count,
                    exc=interrupted,
                    commit_log=commit_log,
                    metrics=metrics,
                )
            )
            async with async_session_context() as session:
                await session.execute(
                    update(PendingFile)
                    .where(PendingFile.file_id == file_id)
                    .values(
                        status=1,
                        error_message=None,
                        retry_count=PendingFile.retry_count + 1,
                        updated_at=datetime.now(),
                    )
                )
                await session.commit()
            raise
        except RedisError as exc:
            # Redis owns distributed concurrency and the per-file fencing lock.
            # A transient Redis outage is infrastructure backpressure, not a
            # permanent document failure. Roll back this attempt and put the
            # source file back into the queue for the still-running worker.
            interrupted = RuntimeError(f"Redis 暂时不可用，任务已安全重新排队: {exc}")
            logger.warning(
                "[self-ingestion] Redis 异常，入库已回队重试: file_id=%s error=%s",
                file_id,
                exc,
            )
            _log_stage(states.STAGE_FAILED, "redis_retry", detail=str(exc)[:300])
            await self._rollback_failure(
                file_id=file_id,
                job_id=job_id,
                doc_id=doc_id,
                kb_code=kb_code,
                commit_version=commit_version,
                vector_count=vector_count,
                exc=interrupted,
                commit_log=commit_log,
                metrics=metrics,
            )
            async with async_session_context() as session:
                await session.execute(
                    update(PendingFile)
                    .where(PendingFile.file_id == file_id)
                    .values(
                        status=1,
                        error_message=None,
                        retry_count=PendingFile.retry_count + 1,
                        updated_at=datetime.now(),
                    )
                )
                await session.commit()
            return False
        except Exception as exc:
            logger.error("[self-ingestion] 入库失败: file_id=%s error=%s", file_id, exc, exc_info=True)
            _log_stage(states.STAGE_FAILED, "failed", detail=str(exc)[:300])
            try:
                from backend.rag.ingestion.metrics import (
                    record_error,
                    record_job_status,
                )
                record_error(kb_code or "unknown", states.STAGE_FAILED)
                record_job_status(kb_code or "unknown", states.JOB_FAILED)
            except Exception:
                pass
            await self._rollback_failure(
                file_id=file_id, job_id=job_id, doc_id=doc_id, kb_code=kb_code,
                commit_version=commit_version, vector_count=vector_count, exc=exc,
                commit_log=commit_log, metrics=metrics,
            )
            return False
        finally:
            if downloaded_path is not None:
                await asyncio.to_thread(downloaded_path.unlink, missing_ok=True)

    @staticmethod
    def _build_image_assets(
        parsed,
        *,
        content: bytes,
        file_name: str,
        file_ext: str,
        mime_type: str,
        embedded_assets: dict[str, ImageAsset] | None = None,
    ) -> dict[str, ImageAsset]:
        """Attach raw image bytes to image blocks for VLM entity/relation extraction."""
        if parsed is None or not getattr(parsed, "blocks", None):
            return dict(embedded_assets or {})
        ext = (file_ext or Path(file_name or "").suffix).lower()
        mime = (mime_type or "").lower().strip()
        is_standalone_image = ext in IMAGE_EXTENSIONS or mime in IMAGE_MIME_TYPES
        mime_map = {
            ".png": "image/png",
            ".jpg": "image/jpeg",
            ".jpeg": "image/jpeg",
            ".gif": "image/gif",
            ".webp": "image/webp",
            ".bmp": "image/bmp",
            ".tiff": "image/tiff",
        }
        default_mime = mime_map.get(ext, mime or "image/png")
        assets: dict[str, ImageAsset] = dict(embedded_assets or {})
        for block in parsed.blocks:
            if block.block_type != "image":
                continue
            if block.block_id in assets:
                continue
            if is_standalone_image and content:
                assets[block.block_id] = ImageAsset(
                    block_id=block.block_id,
                    content=content,
                    content_type=default_mime,
                    file_name=file_name or block.original_file_name,
                )
        return assets

    async def _start_job(self, session, job_id: str, pending_file: PendingFile) -> None:
        now = datetime.now()
        job = SelfIngestJob(
            job_id=job_id,
            kb_id=pending_file.kb_id,
            file_id=pending_file.file_id,
            source_hash=pending_file.source_hash,
            status=states.JOB_PROCESSING,
            stage=states.STAGE_STARTED,
            started_at=now,
        )
        session.add(job)
        await session.execute(
            update(PendingFile)
            .where(PendingFile.file_id == pending_file.file_id)
            .values(
                ingestion_backend=SELF_BACKEND,
                ingest_job_id=job_id,
                last_ingest_started_at=now,
                error_message=None,
            )
        )

    async def _mark_job(
        self, session, job_id: str, *, status: str, stage: str,
        doc_id: str | None = None, source_hash: str | None = None,
        error_message: str | None = None,
        extra: dict | None = None,
    ) -> None:
        values: dict = {
            "status": status,
            "stage": stage,
            "updated_at": datetime.now(),
        }
        if doc_id is not None:
            values["doc_id"] = doc_id
        if source_hash is not None:
            values["source_hash"] = source_hash
        if error_message is not None:
            values["error_message"] = error_message
        if status in (states.JOB_COMPLETED, states.JOB_FAILED, states.JOB_PARTIAL_SUCCESS):
            values["finished_at"] = datetime.now()
        if extra is not None:
            values["extra_json"] = extra
        await session.execute(
            update(SelfIngestJob).where(SelfIngestJob.job_id == job_id).values(**values)
        )

    async def _finalize_job(
        self, session, job_id: str, *, status: str, stage: str,
        doc_id: str | None = None, source_hash: str | None = None,
        error_message: str | None = None,
        commit_log: list[dict] | None = None,
        metrics: dict | None = None,
        extra: dict | None = None,
    ) -> None:
        """终态写入：status/stage/finished_at + commit_log_json + metrics_json + extra_json。"""
        values: dict = {
            "status": status,
            "stage": stage,
            "updated_at": datetime.now(),
        }
        if doc_id is not None:
            values["doc_id"] = doc_id
        if source_hash is not None:
            values["source_hash"] = source_hash
        if error_message is not None:
            values["error_message"] = error_message[:500]
        if status in (states.JOB_COMPLETED, states.JOB_FAILED, states.JOB_PARTIAL_SUCCESS):
            values["finished_at"] = datetime.now()
        if commit_log is not None:
            values["commit_log_json"] = commit_log[-100:]
        if metrics is not None:
            values["metrics_json"] = metrics
        if extra is not None:
            values["extra_json"] = extra
        await session.execute(
            update(SelfIngestJob).where(SelfIngestJob.job_id == job_id).values(**values)
        )

    async def _write_commit_log_redis(
        self, kb_code: str, job_id: str, commit_log: list[dict], final_status: str,
    ) -> None:
        """将 commit_log 写入 Redis 审计 key（方案 §17 commit_log:{job_id}）。"""
        if not settings.redis_host or not commit_log:
            return
        from backend.db.redis import get_redis
        client = get_redis()
        try:
            key = f"{kb_code}:v2:commit_log:{job_id}"
            payload = {"job_id": job_id, "final_status": final_status, "entries": commit_log[-100:]}
            import json as _json
            await client.set(key, _json.dumps(payload, ensure_ascii=False), ex=86400 * 7)
        except Exception as exc:
            logger.debug("[self-ingestion] commit_log Redis 写入失败: %s", exc)
        finally:
            await client.aclose()

    async def _register_doc(
        self,
        session,
        *,
        pending_file: PendingFile,
        job_id: str,
        doc_id: str,
        source_hash: str,
    ) -> tuple[str, int]:
        """事务内 SELECT FOR UPDATE 判定重复与版本（S4）。

        Returns (action, commit_version):
            action = "skip_existing" | "reingest" | "new"
        """
        existing = (
            await session.execute(
                select(KnowledgeBaseDocument)
                .where(
                    KnowledgeBaseDocument.kb_id == pending_file.kb_id,
                    KnowledgeBaseDocument.doc_id == doc_id,
                )
                .with_for_update()
            )
        ).scalar_one_or_none()

        now = datetime.now()
        if existing is None:
            session.add(
                KnowledgeBaseDocument(
                    kb_id=pending_file.kb_id,
                    doc_id=doc_id,
                    file_name=pending_file.file_name,
                    file_path=pending_file.minio_path,
                    file_size=pending_file.file_size,
                    source_hash=source_hash,
                    object_bucket=pending_file.object_bucket,
                    object_key=pending_file.object_key,
                    object_storage_uri=pending_file.original_uri or pending_file.minio_path,
                    mime_type=pending_file.mime_type or pending_file.file_type,
                    file_ext=pending_file.file_ext,
                    document_date=pending_file.document_date,
                    category_path=pending_file.category_path,
                    ingestion_backend=SELF_BACKEND,
                    doc_status=states.DOC_STAGED,
                    commit_version=1,
                    latest_job_id=job_id,
                    storage_status=0,
                    uploaded_at=pending_file.created_at,
                    ingest_started_at=now,
                    created_at=now,
                    updated_at=now,
                )
            )
            return "new", 1

        existing_version = int(existing.commit_version or 0)
        existing_status = existing.doc_status
        existing_hash = existing.source_hash

        # 只有已经 visible/active 且内容相同才跳过。staged/failed 说明上次没有
        # 完成可见提交，必须升版重放，避免撞上残留的 chunk/block 唯一键。
        if (
            existing_hash == source_hash
            and existing_status in states.RETRIEVAL_VISIBLE_DOC_STATUSES
            and int(existing.storage_status or 0) == 1
        ):
            return "skip_existing", max(existing_version, 1)

        if (
            existing_status in states.RETRIEVAL_VISIBLE_DOC_STATUSES
            and existing.latest_job_id
        ):
            active_status = (
                await session.execute(
                    select(SelfIngestJob.status).where(
                        SelfIngestJob.job_id == existing.latest_job_id
                    )
                )
            ).scalar_one_or_none()
            if active_status in states.JOB_ACTIVE_STATUSES:
                raise RuntimeError(
                    f"同一文档已有入库任务进行中: doc_id={doc_id} job_id={existing.latest_job_id}"
                )

        next_version = max(existing_version, 0) + 1
        existing.latest_job_id = job_id
        existing.commit_version = next_version
        existing.ingest_started_at = now
        existing.updated_at = now
        if existing_status not in states.RETRIEVAL_VISIBLE_DOC_STATUSES:
            existing.file_name = pending_file.file_name
            existing.file_path = pending_file.minio_path
            existing.file_size = pending_file.file_size
            existing.source_hash = source_hash
            existing.object_bucket = pending_file.object_bucket
            existing.object_key = pending_file.object_key
            existing.object_storage_uri = pending_file.original_uri or pending_file.minio_path
            existing.mime_type = pending_file.mime_type or pending_file.file_type
            existing.file_ext = pending_file.file_ext
            existing.document_date = pending_file.document_date
            existing.category_path = pending_file.category_path
            existing.ingestion_backend = SELF_BACKEND
            existing.doc_status = states.DOC_STAGED
            existing.storage_status = 0
        return "reingest", next_version

    async def _rollback_failure(
        self,
        *,
        file_id: str,
        job_id: str,
        doc_id: str | None,
        kb_code: str | None,
        commit_version: int,
        vector_count: int,
        exc: Exception,
        commit_log: list[dict] | None = None,
        metrics: dict | None = None,
    ) -> None:
        """失败回滚：新版本 chunks -> failed，旧版本保持原状（S1）。"""
        error_msg = str(exc)
        now = datetime.now()

        # Milvus：仅删除本次 commit_version 的向量，不动旧版本
        if doc_id and kb_code and vector_count:
            try:
                await MilvusChunkWriter(kb_code).delete_chunks_by_version(doc_id, commit_version)
            except Exception as cleanup_exc:
                logger.warning(
                    "[self-ingestion] rollback Milvus 清理失败: doc=%s ver=%s error=%s",
                    doc_id, commit_version, cleanup_exc,
                )

        try:
            async with async_session_context() as session:
                if doc_id:
                    # 新版本 chunks -> failed
                    await session.execute(
                        update(SelfIngestChunk)
                        .where(
                            SelfIngestChunk.doc_id == doc_id,
                            SelfIngestChunk.commit_version == commit_version,
                            SelfIngestChunk.doc_status == states.DOC_STAGED,
                        )
                        .values(
                            is_active=False,
                            is_visible=False,
                            doc_status=states.DOC_FAILED,
                            embedding_status=states.EMBEDDING_FAILED,
                            updated_at=now,
                        )
                    )
                    # doc -> failed（仅当当前还是 staged，避免覆盖已 visible 的旧状态）
                    await session.execute(
                        update(KnowledgeBaseDocument)
                        .where(
                            KnowledgeBaseDocument.doc_id == doc_id,
                            KnowledgeBaseDocument.doc_status == states.DOC_STAGED,
                        )
                        .values(
                            doc_status=states.DOC_FAILED,
                            storage_status=0,
                            updated_at=now,
                        )
                    )
                await session.execute(
                    update(PendingFile)
                    .where(PendingFile.file_id == file_id)
                    .values(
                        status=3,
                        error_message=error_msg[:500],
                        last_ingest_finished_at=now,
                    )
                )
                await self._finalize_job(
                    session, job_id, status=states.JOB_FAILED,
                    stage=states.STAGE_FAILED, doc_id=doc_id,
                    error_message=error_msg[:500],
                    commit_log=commit_log, metrics=metrics,
                )
                await session.commit()
        except Exception as db_err:
            logger.error(
                "[self-ingestion] rollback 写入失败状态时再次出错: file_id=%s error=%s",
                file_id, db_err, exc_info=True,
            )

    async def _refresh_kb_stats(self, session, kb: KnowledgeBase) -> None:
        """重算知识库聚合统计（doc_count / chunk_count / entity/relation_count）。"""
        await refresh_knowledge_base_stats(session, kb)

    async def _extract_text(
        self,
        pending_file: PendingFile,
        content: bytes,
        file_id: str,
    ) -> str:
        ext = (pending_file.file_ext or Path(pending_file.file_name or "").suffix).lower()
        mime = (pending_file.mime_type or pending_file.file_type or "").lower()
        if ext in PDF_EXTENSIONS or mime in PDF_MIME_TYPES:
            return await self._parse_pdf_to_markdown(pending_file, content, file_id)
        if ext in DOCX_EXTENSIONS or mime in DOCX_MIME_TYPES:
            return await asyncio.to_thread(
                self._parse_docx_to_markdown,
                pending_file,
                content,
            )
        if ext not in TEXT_EXTENSIONS and mime not in TEXT_MIME_TYPES:
            raise SelfIngestionNotReadyError(
                "self_ingestion phase-1 当前支持 md/txt/pdf/docx；"
                f"file={pending_file.file_name}, ext={ext}, mime={mime}"
            )
        for encoding in ("utf-8", "utf-8-sig", "gb18030"):
            try:
                return content.decode(encoding)
            except UnicodeDecodeError:
                continue
        return content.decode("utf-8", errors="replace")

    def _parse_docx_to_markdown(self, pending_file: PendingFile, content: bytes) -> str:
        try:
            from docx import Document
        except Exception as exc:
            raise SelfIngestionNotReadyError(
                "解析 docx 需要 python-docx 依赖，请先安装或切回 legacy_lightrag"
            ) from exc

        document = Document(BytesIO(content))
        parts: list[str] = []
        for paragraph in document.paragraphs:
            text = (paragraph.text or "").strip()
            if text:
                parts.append(text)

        for table in document.tables:
            rows: list[list[str]] = []
            for row in table.rows:
                cells = [
                    " ".join((cell.text or "").split())
                    for cell in row.cells
                ]
                if any(cells):
                    rows.append(cells)
            if not rows:
                continue
            width = max(len(row) for row in rows)
            normalized_rows = [row + [""] * (width - len(row)) for row in rows]
            header = normalized_rows[0]
            table_lines = [
                "| " + " | ".join(header) + " |",
                "| " + " | ".join(["---"] * width) + " |",
            ]
            for row in normalized_rows[1:]:
                table_lines.append("| " + " | ".join(row) + " |")
            parts.append("\n".join(table_lines))

        markdown = "\n\n".join(parts).strip()
        if not markdown:
            raise RuntimeError(f"DOCX 解析结果为空: {pending_file.file_name}")
        return markdown

    async def _parse_pdf_to_markdown(
        self,
        pending_file: PendingFile,
        content: bytes,
        file_id: str,
    ) -> str:
        from backend.core.tools.ocr.mineru_ocr import MinerUOCR, OCRConfig
        from backend.rag.ingestion.concurrency import mineru_concurrency, semaphore

        cache_dir = Path(__file__).resolve().parents[2] / "data" / "cache" / "self_ingestion"
        await asyncio.to_thread(cache_dir.mkdir, parents=True, exist_ok=True)
        suffix = Path(pending_file.file_name or "").suffix or ".pdf"
        safe_name = "".join(ch if ch.isalnum() or ch in ("-", "_") else "_" for ch in Path(pending_file.file_name or "document").stem)
        cache_path = cache_dir / f"{file_id}_{safe_name}{suffix}"
        await asyncio.to_thread(cache_path.write_bytes, content)

        config = OCRConfig(
            backend=settings.mineru_backend,
            lang=settings.mineru_lang,
            method=settings.mineru_method,
            formula_enable=settings.mineru_formula_enable,
            table_enable=settings.mineru_table_enable,
        )
        mineru = MinerUOCR(api_url=settings.mineru_api_url, config=config)
        try:
            logger.info("[self-ingestion] 提交 MinerU PDF 解析: %s", pending_file.file_name)
            mineru_sem = await semaphore("self_ingestion_mineru", mineru_concurrency())
            mineru_timeout = int(
                getattr(settings, "self_ingestion_mineru_timeout", 3600) or 3600
            )
            async with mineru_sem:
                task_id = await mineru.submit_task_async(file_path=str(cache_path), config=config)
                result = await mineru.wait_for_completion_async(
                    task_id=task_id,
                    timeout=mineru_timeout,
                    poll_interval=5,
                )
            if result.status.value != "completed":
                error = result.error_message or result.status.value
                raise RuntimeError(f"MinerU 解析失败: {error}")
            md_result = await mineru.get_task_data_async(
                task_id,
                include_fields=["md", "markdown"],
            )
            markdown = md_result.markdown or ""
            if not markdown.strip():
                raise RuntimeError("MinerU 返回的 markdown 为空")
            return markdown
        finally:
            try:
                await mineru.close()
            finally:
                try:
                    await asyncio.to_thread(cache_path.unlink)
                except FileNotFoundError:
                    pass
