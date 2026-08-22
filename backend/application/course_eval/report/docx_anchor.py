"""docx 坐标锚点填充原语（M-07）。

纯函数模块：只依赖 python-docx，不 import 业务实体、不碰数据库。
四个坑的解药（方案 §6.7.1）：
- 坑 1 合并单元格静默摧毁标签 → write_guarded 双重断言；
- 坑 2 术语不一致导致勾选静默失败 → tick_checkbox 返回 bool，调用方必须检查；
- 坑 3 未使用列残留预置符 → clear_unused_columns；
- 坑 4 合并单元格重复计数 → iter_unique_cells 按 id(cell._tc) 去重。
"""

from __future__ import annotations

from dataclasses import dataclass
from typing import Iterable, Iterator

from docx.table import Table

from backend.application.exceptions import ApplicationError


class TemplateAnchorMismatch(ApplicationError):
    """守卫标签断言失败或目标格与守卫格是同一个合并单元格（坑 1）。"""

    status_code = 500
    public_detail = "报告模板锚点与模板结构不符，渲染中止"
    error_code = "course_eval_template_anchor_mismatch"


@dataclass(frozen=True)
class CellAnchor:
    """值格锚点。每个值格必须登记守卫标签及其坐标（坑 1 的硬规则）。"""

    name: str
    row: int
    col: int
    guard: str
    guard_row: int
    guard_col: int
    preset: str = ""


def iter_unique_cells(table: Table) -> Iterator[tuple[object, int, object]]:
    """按 id(cell._tc) 去重遍历（坑 4）。合并单元格只返回一次。"""
    seen: set[int] = set()
    for row in table.rows:
        for ci, cell in enumerate(row.cells):
            key = id(cell._tc)
            if key in seen:
                continue
            seen.add(key)
            yield row, ci, cell


def build_header_index(table: Table, header_row: int = 0) -> dict[str, int]:
    """按表头文本建立 序号→首个匹配列索引，解决合并单元格导致的列错位（§1.4）。"""
    index: dict[str, int] = {}
    for ci, cell in enumerate(table.rows[header_row].cells):
        key = cell.text.strip()
        if key and key not in index:
            index[key] = ci
    return index


def set_cell_text(cell: object, text: str, *, keep_style: bool = True) -> None:
    """整体覆盖单元格文本，保留首个 run 的字体样式。"""
    para = cell.paragraphs[0]
    if keep_style and para.runs:
        para.runs[0].text = text
        for extra in para.runs[1:]:
            extra.text = ""
    else:
        para.text = text
    for extra_para in cell.paragraphs[1:]:
        extra_para._element.getparent().remove(extra_para._element)


def write_guarded(table: Table, anchor: CellAnchor, text: str) -> None:
    """带守卫的写入：防坑 1。

    写前双重断言：守卫标签仍在原位 + 目标格与守卫格不是同一合并单元格。
    任一失败即抛 TemplateAnchorMismatch，中止渲染。
    """
    guard_cell = table.rows[anchor.guard_row].cells[anchor.guard_col]
    if guard_cell.text.strip() != anchor.guard:
        raise TemplateAnchorMismatch(
            f"{anchor.name}: 守卫标签期望 {anchor.guard!r}，实际 {guard_cell.text.strip()!r}"
        )
    target = table.rows[anchor.row].cells[anchor.col]
    if id(target._tc) == id(guard_cell._tc):
        raise TemplateAnchorMismatch(
            f"{anchor.name}: 目标格与守卫标签格是同一个合并单元格，锚点坐标错误"
        )
    set_cell_text(target, text)


def tick_checkbox(cell: object, option_label: str, *, mark: str = "☑") -> bool:
    """把 '□<option_label>' 替换为 '☑<option_label>'。未命中返回 False，调用方必须检查。"""
    for para in cell.paragraphs:
        for run in para.runs:
            if f"□{option_label}" in run.text:
                run.text = run.text.replace(f"□{option_label}", f"{mark}{option_label}")
                return True
    # 段落整体文本也可能承载选项（无 run 分段时）
    if f"□{option_label}" in cell.text:
        para = cell.paragraphs[0]
        para.text = para.text.replace(f"□{option_label}", f"{mark}{option_label}")
        return True
    return False


def clear_unused_columns(
    table: Table,
    *,
    rows: Iterable[int],
    used: set[int],
    all_cols: Iterable[int],
) -> None:
    """清空未使用列（坑 3）。all_cols - used 的每个单元格置空。

    不要只清空「序号大于大题数」的列——表 1 的列索引是跳号的，used 集合才是唯一依据。
    """
    unused = set(all_cols) - set(used)
    for ri in rows:
        for ci in unused:
            if ci < len(table.columns):
                set_cell_text(table.rows[ri].cells[ci], "")


