"""试卷分析表渲染器（M-07）。

输入全是值对象（snapshot 等），不碰数据库，返回 docx 字节。
覆盖学院模板的四张表：
- 表 0 基本情况（课程/学时/学分/考核方式/年级/专业/班号/人数/时间）
- 表 1 题型及分值分布（题型/小题个数/分值/总分 + 试题总个数/合计总分）
- 表 2 成绩分布（六段百分比 + 平均/及格率/最高/最低）
- 表 3 试卷得分率（满分/平均得分/得分率，未用列清预置 %）
定性文字段（不及格原因/结论/改进措施）留白由教师撰写。
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

from docx import Document

from backend.application.course_eval.report.docx_anchor import (
    CellAnchor,
    TemplateAnchorMismatch,
    build_header_index,
    clear_unused_columns,
    set_cell_text,
    tick_checkbox,
    write_guarded,
)
from backend.application.course_eval.stats_engine import (
    NF_AVG_2DP,
    NF_INT,
    NF_PCT_1DP,
    Band,
    BandResult,
    fmt,
)

_CN_NUMERALS = {i + 1: c for i, c in enumerate("一二三四五六七八九十")}
_RE_PATTERN_TERM = re.compile(r"(\d{1,3}(?:\.\d)?)\s*分?\s*[×xX*✕]\s*(\d{1,2})")


@dataclass(frozen=True)
class TaskView:
    """渲染所需的学期任务值对象（避免渲染器依赖 ORM）。"""

    course_name: str
    grade: str = ""
    major: str = ""
    class_no: str = ""
    teacher_name: str = ""
    analyzer_name: str = ""
    hours: str = ""
    credit: str = ""
    exam_mode: str = ""
    course_nature: str = ""
    exam_date: str = ""  # 已格式化（"2024 年 12 月 27 日"）
    exam_duration: str = ""  # 已格式化（"120 分钟"）
    expected_students: str = ""


@dataclass(frozen=True)
class SectionView:
    """表 1/表 3 渲染用的大题值对象。"""

    report_order_no: int
    question_type: str = ""
    score_pattern: str = ""  # "2分×10" / "10分×2+15分×2"
    total_score: Decimal | None = None
    item_count: int | None = None
    avg_score: Decimal | None = None
    rate: Decimal | None = None  # 百分点


@dataclass
class RenderResult:
    docx_bytes: bytes = b""
    warnings: list[str] = field(default_factory=list)


# 表 2「成绩分布表」锚点（§6.7.3 实测，2024-2025 版模板）
DISTRIBUTION_TABLE_INDEX = 2
_DIST_AVERAGE = CellAnchor(
    name="distribution.average",
    row=2,
    col=6,
    guard="平均成绩",
    guard_row=1,
    guard_col=6,
)
_DIST_PASS_RATE = CellAnchor(
    name="distribution.pass_rate",
    row=2,
    col=7,
    guard="及格率",
    guard_row=1,
    guard_col=7,
)
_DIST_MAX = CellAnchor(
    name="distribution.max",
    row=2,
    col=8,
    guard="最高分",
    guard_row=1,
    guard_col=8,
)
_DIST_MIN = CellAnchor(
    name="distribution.min",
    row=2,
    col=9,
    guard="最低分",
    guard_row=1,
    guard_col=9,
)


def _format_average(value, spec=NF_AVG_2DP) -> str:
    """平均成绩小数位由模板 number_format_json 配置；默认 2 位。"""
    return fmt(value, spec)


def render_distribution(
    docx_bytes: bytes,
    *,
    final_score_stats: BandResult,
    task: TaskView | None = None,
) -> bytes:
    """渲染表 2 成绩分布（六个分段百分比 + 平均成绩/及格率/最高/最低）。

    c0~c5 与 c7 预置了 '%'，必须整体覆盖而非追加（§6.7.3 表 2）。
    """
    doc = Document(io.BytesIO(docx_bytes))
    dist = doc.tables[DISTRIBUTION_TABLE_INDEX]

    # 六个分段百分比：c0~c5（预置 %，整体覆盖）
    for bi, band in enumerate(final_score_stats.bands[:6]):
        set_cell_text(dist.rows[2].cells[bi], f"{band.pct_str}")

    write_guarded(dist, _DIST_AVERAGE, _format_average(final_score_stats.average))
    write_guarded(dist, _DIST_PASS_RATE, fmt(final_score_stats.pass_rate, NF_PCT_1DP))
    if final_score_stats.max_score is not None:
        write_guarded(dist, _DIST_MAX, fmt(final_score_stats.max_score, NF_INT))
    if final_score_stats.min_score is not None:
        write_guarded(dist, _DIST_MIN, fmt(final_score_stats.min_score, NF_INT))

    return _to_bytes(doc)


def _to_bytes(doc: Document) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── 表 0 基本情况 ──

# 锚点坐标来自真实模板实测（§6.7.3；合成模板与其同构）
_BASIC_ANCHORS: tuple[tuple[str, CellAnchor], ...] = (
    ("course_name", CellAnchor("basic.course_name", 0, 1, "课程名称", 0, 0)),
    ("hours", CellAnchor("basic.hours", 0, 5, "学时", 0, 4)),
    ("credit", CellAnchor("basic.credit", 0, 7, "学分", 0, 6)),
    ("exam_mode", CellAnchor("basic.exam_mode", 0, 11, "考核方式", 0, 9)),
    ("exam_duration", CellAnchor("basic.exam_duration", 1, 7, "题量", 1, 6)),
    ("exam_date", CellAnchor("basic.exam_date", 1, 11, "考试时间", 1, 9)),
    ("grade", CellAnchor("basic.grade", 2, 2, "参加考试年级", 2, 0)),
    ("major", CellAnchor("basic.major", 2, 4, "专业", 2, 3)),
    ("class_no", CellAnchor("basic.class_no", 3, 1, "教学班号", 3, 0)),
    ("expected_students", CellAnchor("basic.students", 3, 10, "考试总人数", 3, 8)),
    ("teacher_name", CellAnchor("basic.teacher", 4, 1, "任课教师", 4, 0)),
    ("analyzer_name", CellAnchor("basic.analyzer", 4, 10, "试卷分析人", 4, 8)),
)

# 业务口径 → 模板复选框选项的措辞映射（坑 2：术语不一致）
_NATURE_TERMS: dict[str, tuple[str, ...]] = {
    "学科基础课程": ("学科基础课", "专业基础课"),
    "学科基础课": ("学科基础课", "专业基础课"),
    "专业基础课程": ("专业基础课",),
    "专业基础课": ("专业基础课",),
    "公共基础课程": ("公共基础课", "通识通修课"),
    "专业课": ("专业课", "专业方向课"),
}


def render_basic_info(doc: Document, task: TaskView, warnings: list[str]) -> None:
    """表 0：有值才写，空值留白由教师补。"""
    table = doc.tables[0]
    values = {
        "course_name": task.course_name,
        "hours": task.hours,
        "credit": task.credit,
        "exam_mode": task.exam_mode,
        "exam_duration": task.exam_duration,
        "exam_date": task.exam_date,
        "grade": task.grade,
        "major": task.major,
        "class_no": task.class_no,
        "expected_students": task.expected_students,
        "teacher_name": task.teacher_name,
        "analyzer_name": task.analyzer_name,
    }
    for key, anchor in _BASIC_ANCHORS:
        value = str(values.get(key, "") or "").strip()
        if value:
            write_guarded(table, anchor, value)

    nature = str(task.course_nature or "").strip()
    if nature:
        cell = table.rows[1].cells[1]
        for option in _NATURE_TERMS.get(nature, (nature,)):
            if tick_checkbox(cell, option):
                break
        else:
            warnings.append(f"课程属性「{nature}」在模板选项中未命中，请手工勾选")


# ── 表 1 题型及分值分布 ──


def _pattern_terms(pattern: str) -> list[tuple[Decimal, int]]:
    terms: list[tuple[Decimal, int]] = []
    for score, count in _RE_PATTERN_TERM.findall(pattern or ""):
        try:
            terms.append((Decimal(score), int(count)))
        except (InvalidOperation, ValueError):
            continue
    return terms


def section_item_count(section: SectionView) -> int | None:
    """小题个数：显式字段优先，否则从分值模式（"2分×10"）推算。"""
    if section.item_count:
        return section.item_count
    terms = _pattern_terms(section.score_pattern)
    return sum(count for _score, count in terms) or None


def section_item_score_text(section: SectionView) -> str:
    """「分  值」行：统一分值写单题分（"2"），混合分值写紧凑模式（"10×2+15×2"）。"""
    terms = _pattern_terms(section.score_pattern)
    if not terms:
        return ""
    if len(terms) == 1:
        return fmt(terms[0][0], NF_INT)
    return "+".join(f"{fmt(s, NF_INT)}×{c}" for s, c in terms)


def render_section_structure(
    doc: Document, sections: list[SectionView], warnings: list[str]
) -> None:
    """表 1：按报告序号列写题型/小题个数/分值/总分，汇总行写总个数与合计总分。"""
    table = doc.tables[1]
    header = build_header_index(table)
    row_labels = ("题型", "小题个数", "分值", "总分")
    for ri, expected in enumerate(row_labels, start=1):
        actual = table.rows[ri].cells[0].text.replace(" ", "").strip()
        if expected not in actual:
            raise TemplateAnchorMismatch(
                f"表 1 第{ri}行标签期望含「{expected}」，实际「{actual}」"
            )

    used_cols: set[int] = set()
    total_items = 0
    items_known = True
    total_full = Decimal(0)
    for section in sections:
        numeral = _CN_NUMERALS.get(section.report_order_no)
        col = header.get(numeral or "")
        if col is None:
            warnings.append(
                f"第{section.report_order_no}题超出模板列（一~十），未写入表 1"
            )
            continue
        used_cols.add(col)
        set_cell_text(table.rows[1].cells[col], section.question_type or "")
        count = section_item_count(section)
        set_cell_text(table.rows[2].cells[col], str(count) if count else "")
        if count:
            total_items += count
        else:
            items_known = False
        set_cell_text(table.rows[3].cells[col], section_item_score_text(section))
        if section.total_score is not None:
            set_cell_text(
                table.rows[4].cells[col], fmt(Decimal(section.total_score), NF_INT)
            )
            total_full += Decimal(section.total_score)

    all_cols = [ci for ci in header.values() if ci > 0]
    clear_unused_columns(table, rows=(1, 2, 3, 4), used=used_cols, all_cols=all_cols)

    summary = table.rows[5]
    if "试题总个数" in summary.cells[0].text.replace(" ", ""):
        if items_known and total_items:
            set_cell_text(summary.cells[1], f"{total_items} 道")
        if total_full:
            set_cell_text(summary.cells[4], f"{fmt(total_full, NF_INT)} 分")
    else:
        warnings.append("表 1 汇总行结构与模板不符，试题总个数/合计总分未写入")


# ── 表 3 试卷得分率 ──


def render_section_rates(
    doc: Document, sections: list[SectionView], warnings: list[str]
) -> None:
    """表 3：满分/平均得分/得分率；未使用列清空预置 '%'（坑 3）。"""
    table = doc.tables[3]
    header = build_header_index(table)
    used_cols: set[int] = set()
    for section in sections:
        numeral = _CN_NUMERALS.get(section.report_order_no)
        col = header.get(numeral or "")
        if col is None:
            warnings.append(
                f"第{section.report_order_no}题超出模板列（一~十），未写入表 3"
            )
            continue
        used_cols.add(col)
        # 空值也要显式清格：得分率行预置了 '%'，留着会渲染出孤立的百分号（坑 3）
        set_cell_text(
            table.rows[1].cells[col],
            fmt(Decimal(section.total_score), NF_INT)
            if section.total_score is not None
            else "",
        )
        set_cell_text(
            table.rows[2].cells[col],
            fmt(Decimal(section.avg_score), NF_AVG_2DP)
            if section.avg_score is not None
            else "",
        )
        set_cell_text(
            table.rows[3].cells[col],
            fmt(Decimal(section.rate), NF_PCT_1DP)
            if section.rate is not None
            else "",
        )
    all_cols = [ci for ci in header.values() if ci > 0]
    clear_unused_columns(table, rows=(1, 2, 3), used=used_cols, all_cols=all_cols)


# ── 定性段落（四/五/六/七）填充 ──

_RE_NUMBERED_LABEL = re.compile(r"^[12]、$")


def _set_paragraph_text(para, text: str) -> None:
    """整段覆盖文本，保留首个 run 样式（与 set_cell_text 同思路）。"""
    if para.runs:
        para.runs[0].text = text
        for extra in para.runs[1:]:
            extra.text = ""
    else:
        para.text = text


def render_qualitative_sections(
    doc: Document,
    *,
    fail_reasons: list[str],
    conclusion: str,
    improvements: list[str],
    other_notes: list[str],
    warnings: list[str],
) -> None:
    """把起草好的定性文字填进模板的四/五/六/七段落。

    - 「四、不及格原因」「六、改进意见」写进「1、」「2、」标签段；
    - 「五、结论」「七、其它说明」写进标题后的第一个空白段；
    - 草稿为空的段落保持留白（教师手写）；标题找不到只告警不中断。
    """
    paras = doc.paragraphs

    def find(prefix: str) -> int:
        for i, para in enumerate(paras):
            if para.text.strip().startswith(prefix):
                return i
        return -1

    bounds = {
        "四": find("四、"),
        "五": find("五、"),
        "六": find("六、"),
        "七": find("七、"),
    }
    sign = next(
        (i for i, p in enumerate(paras) if "签字" in p.text),
        len(paras),
    )
    order = [i for i in (*bounds.values(), sign) if i != -1]

    def next_bound(idx: int) -> int:
        following = [i for i in order if i > idx]
        return min(following) if following else len(paras)

    def fill_numbered(section: str, items: list[str]) -> None:
        idx = bounds[section]
        if idx == -1:
            warnings.append(f"模板缺少「{section}、」段落，相应草稿未写入")
            return
        cursor = 0
        for i in range(idx + 1, next_bound(idx)):
            label = paras[i].text.strip()
            if _RE_NUMBERED_LABEL.match(label):
                if cursor < len(items) and items[cursor]:
                    _set_paragraph_text(paras[i], f"{label}{items[cursor]}")
                cursor += 1

    def fill_blank(section: str, text: str) -> None:
        if not text:
            return
        idx = bounds[section]
        if idx == -1:
            warnings.append(f"模板缺少「{section}、」段落，相应草稿未写入")
            return
        for i in range(idx + 1, next_bound(idx)):
            if not paras[i].text.strip():
                _set_paragraph_text(paras[i], text)
                return
        warnings.append(f"「{section}、」段落下没有可写入的空白行")

    fill_numbered("四", fail_reasons)
    fill_blank("五", conclusion)
    fill_numbered("六", improvements)
    fill_blank("七", "；".join(note for note in other_notes if note))


# ── 快照 → 值对象 ──


def band_result_from_snapshot(stats: dict) -> BandResult:
    """把 latest_snapshot 序列化的 final_score_stats 还原成 BandResult。"""
    bands = [
        Band(
            label=str(b.get("label", "")),
            count=int(b.get("count", 0)),
            pct=Decimal(str(b.get("pct", "0"))),
        )
        for b in stats.get("bands", [])
    ]

    def _dec(key: str) -> Decimal | None:
        raw = stats.get(key)
        if raw in (None, ""):
            return None
        return Decimal(str(raw))

    return BandResult(
        bands=bands,
        average=_dec("average") or Decimal(0),
        pass_rate=_dec("pass_rate"),
        max_score=_dec("max"),
        min_score=_dec("min"),
    )


def render_paper_analysis(
    docx_bytes: bytes,
    *,
    task: TaskView,
    sections: list[SectionView],
    final_score_stats: dict | None,
) -> RenderResult:
    """整表渲染：表 0/1/3 + 表 2（成绩分布）。返回 docx 字节与非致命告警。"""
    result = RenderResult()
    doc = Document(io.BytesIO(docx_bytes))
    render_basic_info(doc, task, result.warnings)
    render_section_structure(doc, sections, result.warnings)
    if final_score_stats:
        stats = band_result_from_snapshot(final_score_stats)
        dist = doc.tables[DISTRIBUTION_TABLE_INDEX]
        for bi, band in enumerate(stats.bands[:6]):
            set_cell_text(dist.rows[2].cells[bi], f"{band.pct_str}")
        write_guarded(dist, _DIST_AVERAGE, _format_average(stats.average))
        if stats.pass_rate is not None:
            write_guarded(dist, _DIST_PASS_RATE, fmt(stats.pass_rate, NF_PCT_1DP))
        if stats.max_score is not None:
            write_guarded(dist, _DIST_MAX, fmt(stats.max_score, NF_INT))
        if stats.min_score is not None:
            write_guarded(dist, _DIST_MIN, fmt(stats.min_score, NF_INT))
    else:
        result.warnings.append("统计快照缺少成绩分布，表 2 留空")
    render_section_rates(doc, sections, result.warnings)
    result.docx_bytes = _to_bytes(doc)
    return result


__all__ = [
    "render_distribution",
    "render_paper_analysis",
    "band_result_from_snapshot",
    "RenderResult",
    "SectionView",
    "TaskView",
]
