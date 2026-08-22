"""附表 1/2/3 渲染器（M-07 / T3-04）。

- 附表 1：学院/年级/抽样人数/专业/任课教师 + 达成度评价结论；
- 附表 2：课程总评成绩分析（等级 5 档人数与百分比）+ 达成度情况（期望值/实际值，
  变行数按目标数克隆）+ 平时表现达成度 + 样本达成度分布；
- 附表 3：学生个体明细（变行数按学生数克隆）+ 末行实际达成度平均值。

输入全是值对象，不碰数据库，返回 docx 字节。
"""

from __future__ import annotations

import copy
import io
from decimal import Decimal

from docx import Document
from docx.table import Table

from backend.application.course_eval.report.docx_anchor import set_cell_text
from backend.application.course_eval.stats_engine import NF_ACH_3DP, fmt

# 附表 2 模板常量（真实模板 §1.4：21 行 × 8 列）
_APPENDIX2 = {
    "table_index": 1,
    "grade_row": 5,  # 人数行
    "pct_row": 6,  # 百分比行
    "goal_header_row": 7,  # 「课程目标/达成度期望值/实际达成度平均值」
    "goal_first_row": 9,  # 第一个目标的实际值行（r9）
    "goal_label_col": 1,
    "goal_expected_col": 6,
    "goal_actual_col": 7,
    "regular_row": 14,  # 平时表现达成度
    "regular_actual_col": 7,
    "dist_first_row": 16,  # 样本达成度分布第一行
    "dist_cols": [3, 4, 5, 6, 7],  # （0,0.6）~ [0.9,1]
}

# 附表 3 模板常量（真实模板：10 行 × 9 列）
_APPENDIX3 = {
    "table_index": 2,
    "first_data_row": 3,
    "last_data_row": 8,
    "name_col": 1,
    "regular_col": 2,
    "goal_cols": [3, 4, 5, 6, 7, 8],
    "avg_row": 9,
    "avg_goal_cols": [3, 4, 5, 6, 7, 8],
}


def _clone_row(table: Table, source_row: int, count: int) -> None:
    """按源行克隆 count 行（变行数子表：附表 2 目标数、附表 3 学生数）。"""
    src_tr = table.rows[source_row]._tr
    for _ in range(count):
        new_tr = copy.deepcopy(src_tr)
        src_tr.addnext(new_tr)


def render_appendix2(
    docx_bytes: bytes,
    *,
    grade_counts: list[int],
    grade_pcts: list[Decimal],
    max_score: Decimal,
    min_score: Decimal,
    average: Decimal,
    goals: list[dict],  # [{goal_code, description, expected, actual}]
    regular_achievement: Decimal,
    achievement_dist: list[dict],  # [{goal_code, bands: [(label, count, pct)]}]
) -> bytes:
    """渲染附表 2。达成度情况变行数按 goals 克隆；样本分布按 goals 克隆。"""
    doc = Document(io.BytesIO(docx_bytes))
    table = doc.tables[_APPENDIX2["table_index"]]

    # 等级人数/百分比（r5 人数、r6 百分比，c3~c7）
    for idx, (count, pct) in enumerate(zip(grade_counts, grade_pcts)):
        set_cell_text(table.rows[_APPENDIX2["grade_row"]].cells[3 + idx], str(count))
        set_cell_text(table.rows[_APPENDIX2["pct_row"]].cells[3 + idx], f"{pct}")

    # 达成度情况：变行数（模板 r9 起有 5 行，实际目标数可不同）
    goal_rows_needed = len(goals)
    template_rows = 5  # r9-r13
    if goal_rows_needed > template_rows:
        _clone_row(
            table,
            _APPENDIX2["goal_first_row"] + template_rows - 1,
            goal_rows_needed - template_rows,
        )
    for i, goal in enumerate(goals):
        row = table.rows[_APPENDIX2["goal_first_row"] + i]
        set_cell_text(
            row.cells[_APPENDIX2["goal_label_col"]],
            f"【{goal['goal_code'].replace('目标', '')}】{goal.get('description', '')}",
        )
        set_cell_text(
            row.cells[_APPENDIX2["goal_expected_col"]], str(goal.get("expected", ""))
        )
        set_cell_text(
            row.cells[_APPENDIX2["goal_actual_col"]],
            fmt(Decimal(goal["actual"]), NF_ACH_3DP),
        )

    # 平时表现达成度
    set_cell_text(
        table.rows[_APPENDIX2["regular_row"]].cells[_APPENDIX2["regular_actual_col"]],
        fmt(regular_achievement, NF_ACH_3DP),
    )

    # 样本达成度分布（变行数）
    if achievement_dist:
        dist_needed = len(achievement_dist)
        dist_template = 5  # r16-r20
        if dist_needed > dist_template:
            _clone_row(
                table,
                _APPENDIX2["dist_first_row"] + dist_template - 1,
                dist_needed - dist_template,
            )
        for i, dist in enumerate(achievement_dist):
            row = table.rows[_APPENDIX2["dist_first_row"] + i]
            set_cell_text(row.cells[1], f"【{dist['goal_code'].replace('目标', '')}】")
            for j, (label, count, pct) in enumerate(dist["bands"]):
                col = _APPENDIX2["dist_cols"][j]
                set_cell_text(row.cells[col], f"{count} 人\n占 {pct} %")

    return _to_bytes(doc)


def render_appendix3(
    docx_bytes: bytes,
    *,
    students: list[dict],  # [{name, regular, goal_achievements: list[Decimal]}]
    goal_count: int,
    averages: list[Decimal],  # 每目标实际达成度平均值
) -> bytes:
    """渲染附表 3 学生个体明细（变行数按学生数克隆）+ 末行平均值。"""
    doc = Document(io.BytesIO(docx_bytes))
    table = doc.tables[_APPENDIX3["table_index"]]

    template_rows = _APPENDIX3["last_data_row"] - _APPENDIX3["first_data_row"] + 1
    if len(students) > template_rows:
        _clone_row(table, _APPENDIX3["last_data_row"], len(students) - template_rows)

    for i, student in enumerate(students):
        row = table.rows[_APPENDIX3["first_data_row"] + i]
        set_cell_text(row.cells[0], str(i + 1))
        set_cell_text(row.cells[_APPENDIX3["name_col"]], student["name"])
        set_cell_text(
            row.cells[_APPENDIX3["regular_col"]],
            (
                f"{Decimal(student['regular']):.2f}"
                if student.get("regular") is not None
                else ""
            ),
        )
        for j, value in enumerate(student["goal_achievements"]):
            if j < len(_APPENDIX3["goal_cols"]):
                set_cell_text(
                    row.cells[_APPENDIX3["goal_cols"][j]],
                    fmt(Decimal(value), NF_ACH_3DP) if value is not None else "",
                )

    # 末行平均值
    avg_row = table.rows[_APPENDIX3["avg_row"]]
    for j, value in enumerate(averages):
        if j < len(_APPENDIX3["avg_goal_cols"]):
            set_cell_text(
                avg_row.cells[_APPENDIX3["avg_goal_cols"][j]],
                fmt(Decimal(value), NF_ACH_3DP),
            )

    return _to_bytes(doc)


def render_appendix1(
    docx_bytes: bytes,
    *,
    college: str,
    grade: str,
    sample_count: int,
    major: str,
    teacher_name: str,
    goals: list[dict],  # [{goal_code, description, expected, actual}]
    conclusion: str = "",
    problems: str = "",
    improvement: str = "",
) -> bytes:
    """渲染附表 1：基本信息 + 课程目标达成度评价结论。"""
    doc = Document(io.BytesIO(docx_bytes))
    table = doc.tables[0]

    # r0: 学院 / 开课年级 / 抽样人数
    set_cell_text(table.rows[0].cells[1], college)
    set_cell_text(table.rows[0].cells[4], grade)
    set_cell_text(table.rows[0].cells[7], str(sample_count))
    # r1: 专业 / 任课教师
    set_cell_text(table.rows[1].cells[1], major)
    set_cell_text(table.rows[1].cells[4], teacher_name)

    # r2: 课程目标（换行列出）
    goal_text = "\n".join(
        f"【{g['goal_code'].replace('目标', '')}】{g['description']}：实际达成 {fmt(Decimal(g['actual']), NF_ACH_3DP)}"
        for g in goals
    )
    set_cell_text(table.rows[2].cells[1], goal_text)

    # r3/r4/r5: 达成度分析结果评价 / 存在的问题及原因分析 / 持续改进
    if conclusion:
        set_cell_text(table.rows[3].cells[1], conclusion)
    if problems:
        set_cell_text(table.rows[4].cells[1], problems)
    if improvement and len(table.rows) > 5:
        set_cell_text(table.rows[5].cells[1], improvement)

    return _to_bytes(doc)


def _to_bytes(doc: Document) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
