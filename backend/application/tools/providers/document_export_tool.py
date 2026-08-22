"""export_document：通用交付物导出工具。

把智能体回答（Markdown）导出为 DOCX 或 PDF 交付物并存储为会话工件，返回
下载/工件元数据。适用于所有 full 工作流的日常问答——此前只有 deep_research
的 `save_deep_research_report` 能生成 DOCX，普通对话的长答案无法交付化；本工具
把该能力通用化，并新增 PDF 导出（LibreOffice headless）。

管道：Markdown --pandoc(gfm)--> DOCX --(可选)LibreOffice--> PDF。
- 图表签名 URL（visualize_excel_data / execute_python 的 `/api/v1/charts/...`）
  在渲染前验签并物化为本地 PNG，让 Pandoc 能把图表嵌进文档；
- 其余图片引用（含 `![](http://...)`、reference 定义、裸 `<img>`）一律剥离，
  防止 Pandoc 在进程内变成 SSRF / 任意本地文件读取原语；
- DOCX 走 MinIO（本地降级）存储，与 `save_deep_research_report` 同一工件通路。
"""

from __future__ import annotations

import re
import subprocess
import tempfile
import uuid
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, urlparse

from langchain.tools import tool

from backend.application.blocking import run_blocking, run_in_fresh_loop
from backend.config.loggers import get_logger
from backend.services.document_conversion import convert_docx_to_pdf
from backend.storage import get_user_minio_bucket_name, upload_bytes_to_minio

logger = get_logger(__name__)

DOCX_MIME_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
PDF_MIME_TYPE = "application/pdf"
_MAX_CONTENT_CHARS = 2_000_000
_PANDOC_TIMEOUT = 45.0

# 模块级别名既是既有 patch 点，也让转换实现只有一处（共享服务）。
_convert_docx_to_pdf_bytes = convert_docx_to_pdf

_CHART_IMAGE_URL_RE = re.compile(
    r"(!\[[^\]]*\]\()(/api/v1/charts/[\w-]+/[\w-]+\?[^)\s]+)(\))"
)


def _bounded_content(content: str) -> str:
    text = content or ""
    if len(text) <= _MAX_CONTENT_CHARS:
        return text
    logger.warning(
        "[ExportDocument] 内容超长（%d），已截断至 %d",
        len(text),
        _MAX_CONTENT_CHARS,
    )
    return text[:_MAX_CONTENT_CHARS] + "\n\n[内容超出长度上限，已截断。]"


def _safe_filename(title: str, extension: str) -> str:
    trimmed = (title or "").strip() or "document"
    cleaned = re.sub(r"[\\/:*?\"<>|\x00-\x1f]+", "-", trimmed)
    cleaned = re.sub(r"[\s-]+", "-", cleaned).strip(".-") or "document"
    for suffix in (".docx", ".pdf", ".md"):
        if cleaned.lower().endswith(suffix):
            cleaned = cleaned[: -len(suffix)]
    return f"{cleaned}.{extension}"


def _materialize_chart_images(markdown_input: str, tmp_path: Path) -> str:
    """把图表签名 URL 物化为本地临时 PNG，让 Pandoc 能嵌入图表。

    嵌入前按 HTTP 端点相同规则验签；验签失败或取图失败保留原引用。
    """
    if "/api/v1/charts/" not in markdown_input:
        return markdown_input

    from backend.services.chart_image_store import (
        fetch_chart,
        verify_chart_signature,
    )

    def _replace(match: re.Match) -> str:
        raw_url = match.group(2)
        try:
            parsed = urlparse(raw_url)
            user_id, chart_id = parsed.path.strip("/").split("/")[-2:]
            query = parse_qs(parsed.query)
            exp = int((query.get("exp") or ["0"])[0] or 0)
            sig = (query.get("sig") or [""])[0]
            if not verify_chart_signature(user_id, chart_id, exp, sig):
                return match.group(0)
            # 本函数总在无事件循环的工作线程中执行（run_blocking）；一次性
            # 循环退出前需清理按循环注册的连接资源。
            png = run_in_fresh_loop(lambda: fetch_chart(user_id, chart_id))
            if not png:
                return match.group(0)
            local = tmp_path / f"chart-{user_id}-{chart_id}.png"
            local.write_bytes(png)
            return f"{match.group(1)}{local.as_posix()}{match.group(3)}"
        except Exception:
            logger.warning(
                "[ExportDocument] 图表本地化失败，保留原引用: %s",
                raw_url,
                exc_info=True,
            )
            return match.group(0)

    return _CHART_IMAGE_URL_RE.sub(_replace, markdown_input)


def _sanitize_image_sources(markdown_input: str, tmp_path: Path) -> str:
    """只允许 tmp_path 内的本地图表图片，剥离其余全部图片引用。

    模型生成的 Markdown 逐字进入 Pandoc；`![](http://169.254.169.254/…)`
    会被 Pandoc 在生成文档时抓取，构成进程内 SSRF / 任意本地文件读取。
    因此只放行刚物化到 tmp_path 的签名图表，其余（内联、reference 定义、
    裸 `<img>`）全部剥离。
    """
    allowed_prefix = tmp_path.resolve().as_posix()

    def _is_local_chart(target: str) -> bool:
        cleaned = str(target or "").strip().strip("<>").strip().strip('"').strip("'")
        if not cleaned:
            return False
        try:
            resolved = Path(cleaned).resolve().as_posix()
        except (OSError, RuntimeError, ValueError):
            return False
        return resolved == allowed_prefix or resolved.startswith(allowed_prefix + "/")

    def _inline(match: re.Match) -> str:
        alt, target = match.group(1), match.group(2)
        if _is_local_chart(target):
            return match.group(0)
        return alt or ""

    sanitized = re.sub(
        r"!\[([^\]]*)\]\(\s*<?([^)\s>]+)>?[^)]*\)",
        _inline,
        markdown_input,
    )

    def _ref_definition(match: re.Match) -> str:
        target = match.group(2)
        return match.group(0) if _is_local_chart(target) else ""

    sanitized = re.sub(
        r"(?m)^[ \t]*\[([^\]]+)\]:[ \t]*<?([^\s>]+)>?.*$",
        _ref_definition,
        sanitized,
    )
    sanitized = re.sub(r"(?is)<img\b[^>]*?>", "", sanitized)
    return sanitized


def _render_docx_with_python(*, title: str, content: str) -> bytes:
    """无 Pandoc 时的降级：用 python-docx 渲染常见 Markdown 构造。"""
    from docx import Document

    doc = Document()
    normalized_title = title.strip() or "文档"
    lines = (content.strip() or "(无内容)").splitlines()
    first_content = next((line for line in lines if line.strip()), "")
    if not re.match(r"^\s*#\s+", first_content):
        doc.add_heading(normalized_title, level=0)

    in_code_block = False
    code_lines: list[str] = []
    for raw_line in lines:
        line = raw_line.rstrip()
        if line.strip().startswith("```"):
            if in_code_block:
                doc.add_paragraph("\n".join(code_lines))
                code_lines = []
            in_code_block = not in_code_block
            continue
        if in_code_block:
            code_lines.append(line)
            continue
        stripped = line.strip()
        if not stripped:
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading:
            level = min(len(heading.group(1)), 4)
            doc.add_heading(heading.group(2).strip(), level=max(1, level - 1))
            continue
        if re.match(r"^[-*+]\s+", stripped):
            doc.add_paragraph(re.sub(r"^[-*+]\s+", "", stripped), style="List Bullet")
            continue
        if re.match(r"^\d+[.)]\s+", stripped):
            doc.add_paragraph(re.sub(r"^\d+[.)]\s+", "", stripped), style="List Number")
            continue
        if re.fullmatch(r"[-*_]{3,}", stripped):
            continue
        doc.add_paragraph(stripped)
    if code_lines:
        doc.add_paragraph("\n".join(code_lines))

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _render_docx_bytes(*, title: str, content: str) -> bytes:
    """Markdown 渲染为 DOCX 字节，Pandoc 失败时回落 python-docx。"""
    normalized_title = title.strip() or "文档"
    normalized_content = _bounded_content(content).strip() or "(无内容)"
    markdown_input = normalized_content
    if not re.match(r"^\s*#\s+", normalized_content):
        markdown_input = f"# {normalized_title}\n\n{normalized_content}"

    try:
        import pypandoc

        pandoc_path = pypandoc.get_pandoc_path()
        with tempfile.TemporaryDirectory(prefix="export-docx-") as tmp_dir:
            tmp_path = Path(tmp_dir)
            input_md = tmp_path / "doc.md"
            output_docx = tmp_path / "doc.docx"
            input_md.write_text(
                _sanitize_image_sources(
                    _materialize_chart_images(markdown_input, tmp_path),
                    tmp_path,
                ),
                encoding="utf-8",
            )
            cmd = [
                pandoc_path,
                str(input_md),
                "-f",
                "gfm",
                "-o",
                str(output_docx),
                "--metadata=lang:zh-CN",
                f"--metadata=title:{normalized_title}",
            ]
            result = subprocess.run(
                cmd, capture_output=True, text=True, timeout=_PANDOC_TIMEOUT
            )
            if result.returncode != 0:
                raise RuntimeError(
                    f"Pandoc conversion failed: {result.stderr.strip() or 'unknown'}"
                )
            return output_docx.read_bytes()
    except Exception as exc:  # noqa: BLE001
        logger.warning("[ExportDocument] Pandoc 不可用，回落 python-docx: %s", exc)
        return _render_docx_with_python(
            title=normalized_title, content=normalized_content
        )


async def _store_export_artifact(
    data: bytes,
    *,
    user_id: str,
    thread_id: str,
    filename: str,
    mime: str,
) -> dict[str, Any]:
    object_name = (
        f"users/{user_id or 'anonymous'}/threads/{thread_id}/exports/"
        f"{uuid.uuid4().hex}-{filename}"
    )
    minio_object: str | None = None
    try:
        minio_object = await upload_bytes_to_minio(
            data,
            object_name,
            content_type=mime,
            bucket_name=get_user_minio_bucket_name(),
        )
    except Exception as exc:  # noqa: BLE001
        logger.warning("[ExportDocument] MinIO 上传失败，尝试本地降级: %s", exc)

    local_path: str | None = None
    if not minio_object:
        local_path = await run_blocking(
            _save_locally,
            thread_id=thread_id,
            filename=filename,
            content=data,
            timeout_seconds=20,
        )
    return {
        "artifact_path": filename,
        "minio_object": minio_object,
        "local_path": local_path,
        "bytes": len(data),
    }


def _save_locally(*, thread_id: str, filename: str, content: bytes) -> str | None:
    try:
        from backend.core.config import get_paths

        paths = get_paths()
        paths.ensure_thread_dirs(thread_id)
        output_path = paths.thread_artifacts_dir(thread_id) / filename
        output_path.write_bytes(content)
        return str(output_path)
    except Exception:
        logger.warning("[ExportDocument] 本地降级保存失败", exc_info=True)
        return None


def get_export_document_tool(runtime_context: Any = None):
    """按运行时身份构造 export_document 工具（provider 工厂）。"""

    user_id = str(getattr(runtime_context, "user_id", "") or "").strip()
    thread_id = str(getattr(runtime_context, "thread_id", "") or "").strip()

    @tool("export_document", parse_docstring=True)
    async def export_document(title: str, content: str, format: str = "docx") -> str:
        """把 Markdown 内容导出为 DOCX 或 PDF 交付文档并保存为会话工件。

        用于用户明确需要把回答/报告导出成可下载文档时。内容用标准 Markdown
        （支持标题、列表、表格、代码块；图表用 visualize_excel_data /
        execute_python 返回的 Markdown 图片链接引用即可自动嵌入）。

        Args:
            title: 文档标题。
            content: 完整的 Markdown 正文。
            format: 导出格式，"docx"（默认）或 "pdf"。
        """
        fmt = str(format or "docx").strip().lower()
        if fmt not in {"docx", "pdf"}:
            return f"【导出失败】不支持的格式 '{format}'，仅支持 docx 或 pdf。"
        if not str(content or "").strip():
            return "【导出失败】content 为空，请提供 Markdown 正文。"
        if not thread_id:
            return "【导出失败】当前运行时缺少会话标识，无法保存工件。"

        docx_bytes = await run_blocking(
            _render_docx_bytes, title=title, content=content, timeout_seconds=60
        )

        if fmt == "pdf":
            pdf_bytes = await run_blocking(
                _convert_docx_to_pdf_bytes, docx_bytes, timeout_seconds=120
            )
            if pdf_bytes is None:
                filename = _safe_filename(title, "docx")
                artifact = await _store_export_artifact(
                    docx_bytes,
                    user_id=user_id,
                    thread_id=thread_id,
                    filename=filename,
                    mime=DOCX_MIME_TYPE,
                )
                return (
                    "【已导出（降级为 DOCX）】PDF 转换不可用，已生成 DOCX 工件："
                    f"{filename}（{artifact['bytes']} 字节），可在右侧工件面板查看/下载。"
                )
            filename = _safe_filename(title, "pdf")
            artifact = await _store_export_artifact(
                pdf_bytes,
                user_id=user_id,
                thread_id=thread_id,
                filename=filename,
                mime=PDF_MIME_TYPE,
            )
            logger.info(
                "[ExportDocument] 已导出 PDF: user=%s thread=%s file=%s bytes=%s",
                user_id,
                thread_id,
                filename,
                artifact["bytes"],
            )
            return (
                f"【已导出 PDF】{filename}（{artifact['bytes']} 字节），"
                "可在右侧工件面板查看/下载。"
            )

        filename = _safe_filename(title, "docx")
        artifact = await _store_export_artifact(
            docx_bytes,
            user_id=user_id,
            thread_id=thread_id,
            filename=filename,
            mime=DOCX_MIME_TYPE,
        )
        logger.info(
            "[ExportDocument] 已导出 DOCX: user=%s thread=%s file=%s bytes=%s",
            user_id,
            thread_id,
            filename,
            artifact["bytes"],
        )
        return (
            f"【已导出 DOCX】{filename}（{artifact['bytes']} 字节），"
            "可在右侧工件面板查看/下载。"
        )

    return export_document
